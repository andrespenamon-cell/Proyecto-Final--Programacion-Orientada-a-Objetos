#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tracker de Hábitos Inteligente — versión Python (Programación Orientada a Objetos)
=================================================================================

Recrea, en código, el libro de Excel del proyecto integrador:

    Excel (hoja)              ->  Python (clase)
    ------------------------      ------------------------------------------
    Configuración             ->  Usuario  (+ parámetros de ModeloAutomaticidad)
    Gestión de Hábitos        ->  Habito   (con Frecuencia y Estado)
    Registro Diario           ->  Registro (con la racha calculada por fecha)
    Estadísticas / Dashboard  ->  Reporte  (clase calculada, no se guarda)
    Cálculo                   ->  ModeloAutomaticidad  A(t) = A_máx·(1 − e^(−kt))
    Modelo vs Real            ->  ModeloVsReal
    (archivo .xlsx)           ->  Repositorio / RepositorioJSON  (persistencia)
    (todo el libro)           ->  Tracker  (fachada que une todas las piezas)

Requisitos: solo Python 3.8 o superior (biblioteca estándar).
Opcionales:  matplotlib  (gráficos PNG)  y  openpyxl  (leer / comparar el Excel).

Uso rápido
----------
    python tracker_habitos.py                 # ventana gráfica (Tkinter)
    python tracker_habitos.py --consola       # versión de texto (menú en la terminal)
    python tracker_habitos.py --demo          # muestra todo con los datos de ejemplo
    python tracker_habitos.py --probar        # autopruebas (deben decir OK)
    python tracker_habitos.py --comparar tracker-habitos-inteligente.xlsx
    python tracker_habitos.py --graficos graficos
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from abc import ABC, abstractmethod
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from enum import Enum
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple, Union

# ======================================================================
# 1. Utilidades de formato (números con coma, fechas, barras de texto)
# ======================================================================


def fnum(valor: float, decimales: int = 1) -> str:
    """Número con coma decimal, como en el Excel en español."""
    return f"{valor:.{decimales}f}".replace(".", ",")


def fpct(fraccion: float, decimales: int = 1) -> str:
    """Fracción 0..1 -> '90,0 %'."""
    return fnum(fraccion * 100, decimales) + " %"


def ftexto_fecha(d: Optional[date]) -> str:
    return d.strftime("%d/%m/%Y") if d else "—"


def parsear_fecha(texto: str) -> date:
    """Acepta AAAA-MM-DD, DD/MM/AAAA y DD-MM-AAAA."""
    texto = texto.strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ErrorValidacion(f"Fecha no válida: «{texto}». Usa AAAA-MM-DD o DD/MM/AAAA.")


def barra(fraccion: float, ancho: int = 20) -> str:
    fraccion = max(0.0, min(1.0, fraccion))
    llenos = int(round(fraccion * ancho))
    return "█" * llenos + "░" * (ancho - llenos)


def tabla(encabezados: Sequence[str], filas: Sequence[Sequence[object]],
          alineacion: Optional[Sequence[str]] = None) -> str:
    """Tabla de texto alineada. alineacion: '<' izquierda, '>' derecha, '^' centro."""
    columnas = len(encabezados)
    alineacion = list(alineacion) if alineacion else ["<"] * columnas
    anchos = [len(str(h)) for h in encabezados]
    for fila in filas:
        for i, celda in enumerate(fila):
            anchos[i] = max(anchos[i], len(str(celda)))

    def linea(celdas: Sequence[object]) -> str:
        return "  ".join(f"{str(c):{a}{w}}" for c, a, w in zip(celdas, alineacion, anchos))

    salida = [linea(encabezados), "  ".join("-" * w for w in anchos)]
    salida += [linea(f) for f in filas]
    return "\n".join(salida)


# ======================================================================
# 2. Errores propios
# ======================================================================


class ErrorTracker(Exception):
    """Error base de la aplicación."""


class ErrorValidacion(ErrorTracker):
    """Dato inválido: fecha mala, hábito duplicado, campo vacío, etc."""


# ======================================================================
# 3. Enumeraciones y listas del sistema (hoja Configuración, columnas D–G)
# ======================================================================

CATEGORIAS = ("Salud", "Estudio", "Productividad", "Bienestar", "Otro")


class Frecuencia(Enum):
    DIARIA = "Diaria"
    SEMANAL = "Semanal"
    MENSUAL = "Mensual"

    @classmethod
    def desde_texto(cls, texto: str) -> "Frecuencia":
        for f in cls:
            if f.value.lower() == texto.strip().lower():
                return f
        raise ErrorValidacion(f"Frecuencia no válida: «{texto}». Opciones: Diaria, Semanal, Mensual.")


class Estado(Enum):
    ACTIVO = "Activo"
    ARCHIVADO = "Archivado"

    @classmethod
    def desde_texto(cls, texto: str) -> "Estado":
        for e in cls:
            if e.value.lower() == texto.strip().lower():
                return e
        raise ErrorValidacion(f"Estado no válido: «{texto}». Opciones: Activo, Archivado.")


# ======================================================================
# 4. Entidades del dominio: Usuario, Habito, Registro
# ======================================================================


class Usuario:
    """Persona dueña de los datos (hoja Configuración)."""

    def __init__(self, nombre: str = "Estudiante", meta_general: str = "",
                 fecha_inicio: Optional[date] = None):
        self.nombre = nombre
        self.meta_general = meta_general
        self.fecha_inicio = fecha_inicio or date.today()

    @property
    def nombre(self) -> str:
        return self._nombre

    @nombre.setter
    def nombre(self, valor: str) -> None:
        valor = (valor or "").strip()
        if not valor:
            raise ErrorValidacion("El nombre del usuario no puede estar vacío.")
        self._nombre = valor

    def configurar(self, nombre: Optional[str] = None, meta_general: Optional[str] = None) -> None:
        if nombre is not None:
            self.nombre = nombre
        if meta_general is not None:
            self.meta_general = meta_general.strip()

    def a_dict(self) -> dict:
        return {"nombre": self.nombre, "meta_general": self.meta_general,
                "fecha_inicio": self.fecha_inicio.isoformat()}

    @classmethod
    def desde_dict(cls, d: dict) -> "Usuario":
        return cls(d.get("nombre", "Estudiante"), d.get("meta_general", ""),
                   date.fromisoformat(d["fecha_inicio"]) if d.get("fecha_inicio") else None)


class Habito:
    """Define QUÉ se mide (molde). No contiene el historial: eso es de Registro."""

    def __init__(self, id_habito: int, nombre: str, categoria: str = "Otro",
                 frecuencia: Frecuencia = Frecuencia.DIARIA,
                 fecha_creacion: Optional[date] = None, estado: Estado = Estado.ACTIVO):
        self._id = int(id_habito)
        self.nombre = nombre
        self.categoria = categoria
        self.frecuencia = frecuencia
        self.fecha_creacion = fecha_creacion or date.today()
        self.estado = estado

    @property
    def id(self) -> int:
        return self._id

    @property
    def nombre(self) -> str:
        return self._nombre

    @nombre.setter
    def nombre(self, valor: str) -> None:
        valor = (valor or "").strip()
        if not valor:
            raise ErrorValidacion("El nombre del hábito no puede estar vacío.")
        self._nombre = valor

    @property
    def categoria(self) -> str:
        return self._categoria

    @categoria.setter
    def categoria(self, valor: str) -> None:
        for c in CATEGORIAS:
            if c.lower() == (valor or "").strip().lower():
                self._categoria = c
                return
        raise ErrorValidacion(f"Categoría no válida: «{valor}». Opciones: {', '.join(CATEGORIAS)}.")

    @property
    def esta_activo(self) -> bool:
        return self.estado is Estado.ACTIVO

    def editar(self, nombre: Optional[str] = None, categoria: Optional[str] = None,
               frecuencia: Optional[Frecuencia] = None) -> None:
        if nombre is not None:
            self.nombre = nombre
        if categoria is not None:
            self.categoria = categoria
        if frecuencia is not None:
            self.frecuencia = frecuencia

    def archivar(self) -> None:
        self.estado = Estado.ARCHIVADO

    def reactivar(self) -> None:
        self.estado = Estado.ACTIVO

    def a_dict(self) -> dict:
        return {"id": self.id, "nombre": self.nombre, "categoria": self.categoria,
                "frecuencia": self.frecuencia.value,
                "fecha_creacion": self.fecha_creacion.isoformat(), "estado": self.estado.value}

    @classmethod
    def desde_dict(cls, d: dict) -> "Habito":
        return cls(d["id"], d["nombre"], d.get("categoria", "Otro"),
                   Frecuencia.desde_texto(d.get("frecuencia", "Diaria")),
                   date.fromisoformat(d["fecha_creacion"]),
                   Estado.desde_texto(d.get("estado", "Activo")))

    def __repr__(self) -> str:
        return f"Habito(id={self.id}, nombre={self.nombre!r}, estado={self.estado.value})"


class Registro:
    """Dato transaccional: un renglón por hábito y día (hoja Registro Diario).

    `racha` es un atributo DERIVADO (columna RachaCorrida): lo calcula el Tracker,
    no se edita a mano y no se guarda en el archivo.
    """

    def __init__(self, fecha: date, id_habito: int, cumplido: bool, nota: str = ""):
        self.fecha = fecha
        self.id_habito = int(id_habito)
        self.cumplido = bool(cumplido)
        self.nota = (nota or "").strip()
        self.racha = 0

    def a_dict(self) -> dict:
        return {"fecha": self.fecha.isoformat(), "id_habito": self.id_habito,
                "cumplido": self.cumplido, "nota": self.nota}

    @classmethod
    def desde_dict(cls, d: dict) -> "Registro":
        return cls(date.fromisoformat(d["fecha"]), d["id_habito"], d["cumplido"], d.get("nota", ""))


# ======================================================================
# 5. Reporte: clase CALCULADA (hojas Estadísticas y Dashboard)
# ======================================================================


@dataclass
class EstadisticaHabito:
    habito: Habito
    racha_actual: int
    mejor_racha: int
    pct_total: float
    pct_7_dias: float
    pct_30_dias: float
    registros: int
    cumplidos: int


class Reporte:
    """Agrega los Registros para producir estadísticas. No guarda datos propios.

    Reglas idénticas a las fórmulas del Excel:
      * racha actual  = RachaCorrida del último registro del hábito (LOOKUP(2,1/…)).
      * mejor racha   = máximo de RachaCorrida (MAXIFS).
      * % total       = cumplidos / registros (COUNTIFS).
      * % 7 y 30 días = igual, pero solo con fechas >= HOY-6 y >= HOY-29.
    """

    def __init__(self, habitos: Sequence[Habito], registros: Sequence[Registro], hoy: date):
        self._habitos = list(habitos)
        self._registros = list(registros)
        self.hoy = hoy

    def _registros_de(self, habito: Habito) -> List[Registro]:
        propios = [r for r in self._registros if r.id_habito == habito.id]
        return sorted(propios, key=lambda r: r.fecha)

    @staticmethod
    def _porcentaje(registros: Sequence[Registro]) -> float:
        return sum(1 for r in registros if r.cumplido) / len(registros) if registros else 0.0

    def estadistica(self, habito: Habito) -> EstadisticaHabito:
        regs = self._registros_de(habito)
        desde7 = self.hoy - timedelta(days=6)
        desde30 = self.hoy - timedelta(days=29)
        return EstadisticaHabito(
            habito=habito,
            racha_actual=regs[-1].racha if regs else 0,
            mejor_racha=max((r.racha for r in regs), default=0),
            pct_total=self._porcentaje(regs),
            pct_7_dias=self._porcentaje([r for r in regs if r.fecha >= desde7]),
            pct_30_dias=self._porcentaje([r for r in regs if r.fecha >= desde30]),
            registros=len(regs),
            cumplidos=sum(1 for r in regs if r.cumplido),
        )

    def estadisticas(self, solo_activos: bool = True) -> List[EstadisticaHabito]:
        return [self.estadistica(h) for h in self._habitos if h.esta_activo or not solo_activos]

    def mejor_racha_global(self, solo_activos: bool = True) -> int:
        return max((e.mejor_racha for e in self.estadisticas(solo_activos)), default=0)

    def promedio_cumplimiento(self, solo_activos: bool = True) -> float:
        est = self.estadisticas(solo_activos)
        return sum(e.pct_total for e in est) / len(est) if est else 0.0


# ======================================================================
# 6. Modelo de Cálculo Diferencial (hoja Cálculo)
# ======================================================================


class ModeloAutomaticidad:
    """Modelo de automaticidad de un hábito:  A(t) = A_máx · (1 − e^(−k·t)).

    Se obtiene de la ecuación diferencial  A'(t) = k·(A_máx − A)  con A(0) = 0.
    k se calibra con el dato de Lally et al. (2010): a los `dias_95` días (mediana 66)
    se alcanza el 95 % de A_máx   ->   k = ln(20) / dias_95.
    """

    def __init__(self, a_max: float = 100.0, dias_95: float = 66.0, epsilon: float = 1.0):
        self.configurar(a_max, dias_95, epsilon)

    def configurar(self, a_max: float, dias_95: float, epsilon: float) -> None:
        if a_max <= 0 or dias_95 <= 0 or epsilon <= 0:
            raise ErrorValidacion("A_máx, los días para el 95 % y ε deben ser mayores que cero.")
        self.a_max = float(a_max)
        self.dias_95 = float(dias_95)
        self.epsilon = float(epsilon)

    @property
    def k(self) -> float:
        return math.log(20) / self.dias_95

    @property
    def t_estrella(self) -> float:
        """Día desde el cual A'(t) < ε  (resuelve A'(t*) = ε). 0 si nunca supera ε."""
        if self.a_max * self.k > self.epsilon:
            return math.log(self.a_max * self.k / self.epsilon) / self.k
        return 0.0

    def automaticidad(self, t: float) -> float:
        """A(t)."""
        return self.a_max * (1 - math.exp(-self.k * t))

    def velocidad(self, t: float) -> float:
        """A'(t) = A_máx·k·e^(−kt): puntos de automaticidad ganados por día."""
        return self.a_max * self.k * math.exp(-self.k * t)

    def aceleracion(self, t: float) -> float:
        """A''(t) = −A_máx·k²·e^(−kt) = −k·A'(t): siempre negativa (el avance frena)."""
        return -self.a_max * self.k ** 2 * math.exp(-self.k * t)

    def tabla(self, hasta: int = 91) -> List[Tuple[int, float, float, float]]:
        return [(t, self.automaticidad(t), self.velocidad(t), self.aceleracion(t))
                for t in range(0, hasta + 1)]

    def a_dict(self) -> dict:
        return {"a_max": self.a_max, "dias_95": self.dias_95, "epsilon": self.epsilon}


@dataclass
class ResultadoVsReal:
    habito: Habito
    primer_registro: date
    ultimo_registro: date
    dias_transcurridos: int            # T
    repeticiones: int                  # n
    ritmo: float                       # n / T
    automaticidad_hoy: float           # A(n)
    repeticiones_faltantes: float      # max(0, dias_95 − n)
    dias_restantes: Optional[int]      # ROUNDUP(faltan / ritmo)
    fecha_consolidacion: Optional[date]


class ModeloVsReal:
    """Lleva las repeticiones REALES de un hábito al modelo (hoja Modelo vs Real).

    Cada día cumplido suma una repetición n; los días fallados no suman, por eso
    la curva real queda por debajo de la ideal. Después del último registro la
    curva se proyecta al ritmo de cumplimiento actual.
    """

    def __init__(self, modelo: ModeloAutomaticidad, habito: Habito, registros: Sequence[Registro]):
        self.modelo = modelo
        self.habito = habito
        self._registros = sorted((r for r in registros if r.id_habito == habito.id),
                                 key=lambda r: r.fecha)

    def resultado(self, ritmo_escenario: Optional[float] = None) -> Optional[ResultadoVsReal]:
        """ritmo_escenario (0..1) permite simular «¿y si de ahora en adelante cumplo X % de los días?»."""
        if not self._registros:
            return None
        primero, ultimo = self._registros[0].fecha, self._registros[-1].fecha
        T = (ultimo - primero).days + 1
        n = sum(1 for r in self._registros if r.cumplido)
        ritmo = n / T if T else 0.0
        faltan = max(0.0, self.modelo.dias_95 - n)
        if ritmo_escenario is not None:
            ritmo = ritmo_escenario
            dias = (0 if faltan == 0 else math.ceil(round(faltan / ritmo, 9))) if ritmo > 0 else None
        elif n > 0:
            # ROUNDUP(faltan / (n/T), 0) con aritmética exacta (sin errores de coma flotante)
            dias = math.ceil(round(faltan * T / n, 9))
        else:
            dias = None
        return ResultadoVsReal(
            habito=self.habito, primer_registro=primero, ultimo_registro=ultimo,
            dias_transcurridos=T, repeticiones=n, ritmo=ritmo,
            automaticidad_hoy=self.modelo.automaticidad(n),
            repeticiones_faltantes=faltan, dias_restantes=dias,
            fecha_consolidacion=(ultimo + timedelta(days=dias)) if dias is not None else None)

    def serie(self, horizonte: int = 90, ritmo_escenario: Optional[float] = None) -> List[dict]:
        """Tabla día a día: cumplimiento acumulado real, n(t), curva real/proyectada e ideal."""
        res = self.resultado(ritmo_escenario)
        if res is None:
            return []
        filas = []
        for t in range(1, horizonte + 1):
            fecha = res.primer_registro + timedelta(days=t - 1)
            if t <= res.dias_transcurridos:
                n_t = float(sum(1 for r in self._registros if r.cumplido and r.fecha <= fecha))
                pct: Optional[float] = n_t / t
                proyeccion = False
            else:
                n_t = res.repeticiones + res.ritmo * (t - res.dias_transcurridos)
                pct = None
                proyeccion = True
            filas.append({"t": t, "fecha": fecha, "pct_real": pct, "n": n_t,
                          "real": self.modelo.automaticidad(n_t),
                          "ideal": self.modelo.automaticidad(t),
                          "a_max": self.modelo.a_max, "proyeccion": proyeccion})
        return filas


# ======================================================================
# 7. Persistencia: Repositorio (abstracto) y sus implementaciones
# ======================================================================


class Repositorio(ABC):
    """Contrato de persistencia. El Tracker solo conoce esta interfaz: se puede
    cambiar JSON por una base de datos sin tocar la lógica del negocio."""

    @abstractmethod
    def cargar(self) -> Optional[dict]:
        ...

    @abstractmethod
    def guardar(self, datos: dict) -> None:
        ...


class RepositorioJSON(Repositorio):
    def __init__(self, ruta: Union[str, Path]):
        self.ruta = Path(ruta)

    def cargar(self) -> Optional[dict]:
        if not self.ruta.exists():
            return None
        try:
            with open(self.ruta, "r", encoding="utf-8") as f:
                return json.load(f)
        except (OSError, json.JSONDecodeError) as e:
            raise ErrorTracker(f"No se pudo leer {self.ruta}: {e}")

    def guardar(self, datos: dict) -> None:
        temporal = self.ruta.with_suffix(self.ruta.suffix + ".tmp")
        with open(temporal, "w", encoding="utf-8") as f:
            json.dump(datos, f, ensure_ascii=False, indent=2)
        temporal.replace(self.ruta)


class RepositorioMemoria(Repositorio):
    """No toca el disco: útil para pruebas y para la demo."""

    def __init__(self) -> None:
        self._datos: Optional[dict] = None

    def cargar(self) -> Optional[dict]:
        return json.loads(json.dumps(self._datos)) if self._datos is not None else None

    def guardar(self, datos: dict) -> None:
        self._datos = json.loads(json.dumps(datos))


# ======================================================================
# 8. Tracker: la fachada que une todas las piezas
# ======================================================================

# Datos de ejemplo idénticos a los del Excel (3 hábitos x 30 días = 90 registros)
_EJEMPLO_INICIO = date(2026, 8, 30)
_EJEMPLO_HABITOS = (
    ("Leer 20 minutos", "Estudio", "SSSSSSSSNSSSSSSSSNSSSSSSSSNSSS"),
    ("Tomar 2 litros de agua", "Salud", "SSSNSSSNSSSNSSSNSSSNSSSNSSSNSS"),
    ("Hacer ejercicio", "Bienestar", "SSNSSNSSNSSNSSNSSNSSNSSNSSSSSS"),
)
_EJEMPLO_NOTAS = {("Leer 20 minutos", 6): "Retomando el hábito",
                  ("Hacer ejercicio", 30): "¡Racha más larga hasta ahora!"}


class Tracker:
    """Orquesta Usuario + Hábitos + Registros + Modelo + Repositorio."""

    def __init__(self, repositorio: Repositorio, reloj: Callable[[], date] = date.today,
                 autoguardar: bool = True):
        self.repositorio = repositorio
        self._reloj = reloj
        self.autoguardar = autoguardar
        self.usuario = Usuario()
        self.modelo = ModeloAutomaticidad()
        self._habitos: List[Habito] = []
        self._registros: List[Registro] = []

    # ---------------------------------------------------------- reloj
    def hoy(self) -> date:
        return self._reloj()

    def fijar_hoy(self, fecha: date) -> None:
        self._reloj = lambda: fecha

    # ---------------------------------------------------------- consultas
    @property
    def habitos(self) -> Tuple[Habito, ...]:
        return tuple(self._habitos)

    @property
    def registros(self) -> Tuple[Registro, ...]:
        return tuple(self._registros)

    def habitos_activos(self) -> List[Habito]:
        return [h for h in self._habitos if h.esta_activo]

    def buscar_habito(self, ref: Union[int, str]) -> Habito:
        """Busca por id (3 o '3') o por nombre (sin distinguir mayúsculas)."""
        if isinstance(ref, int) or (isinstance(ref, str) and ref.strip().isdigit()):
            for h in self._habitos:
                if h.id == int(ref):
                    return h
        else:
            for h in self._habitos:
                if h.nombre.lower() == str(ref).strip().lower():
                    return h
        raise ErrorValidacion(f"No existe el hábito «{ref}».")

    def registros_de(self, habito: Habito) -> List[Registro]:
        return sorted((r for r in self._registros if r.id_habito == habito.id), key=lambda r: r.fecha)

    # ---------------------------------------------------------- gestión de hábitos
    def crear_habito(self, nombre: str, categoria: str = "Otro",
                     frecuencia: Frecuencia = Frecuencia.DIARIA) -> Habito:
        self._validar_nombre_libre(nombre)
        nuevo_id = max((h.id for h in self._habitos), default=0) + 1
        habito = Habito(nuevo_id, nombre, categoria, frecuencia, self.hoy())
        self._habitos.append(habito)
        self._persistir()
        return habito

    def editar_habito(self, ref: Union[int, str], nombre: Optional[str] = None,
                      categoria: Optional[str] = None,
                      frecuencia: Optional[Frecuencia] = None) -> Habito:
        habito = self.buscar_habito(ref)
        if nombre is not None and nombre.strip().lower() != habito.nombre.lower():
            self._validar_nombre_libre(nombre)
        habito.editar(nombre, categoria, frecuencia)   # el historial sigue ligado por id
        self._persistir()
        return habito

    def archivar_habito(self, ref: Union[int, str]) -> Habito:
        habito = self.buscar_habito(ref)
        habito.archivar()                              # el historial se conserva
        self._persistir()
        return habito

    def reactivar_habito(self, ref: Union[int, str]) -> Habito:
        habito = self.buscar_habito(ref)
        habito.reactivar()
        self._persistir()
        return habito

    def _validar_nombre_libre(self, nombre: str) -> None:
        if not (nombre or "").strip():
            raise ErrorValidacion("El nombre del hábito no puede estar vacío.")
        if any(h.nombre.lower() == nombre.strip().lower() for h in self._habitos):
            raise ErrorValidacion(f"Ya existe un hábito llamado «{nombre.strip()}».")

    # ---------------------------------------------------------- registro diario
    def registrar(self, ref: Union[int, str], cumplido: bool, fecha: Optional[date] = None,
                  nota: str = "", reemplazar: bool = False) -> Registro:
        habito = self.buscar_habito(ref)
        if not habito.esta_activo:
            raise ErrorValidacion(f"«{habito.nombre}» está archivado: reactívalo para registrar nuevos días.")
        fecha = fecha or self.hoy()
        if fecha > self.hoy():
            raise ErrorValidacion("No se puede registrar una fecha futura.")
        existente = next((r for r in self._registros
                          if r.id_habito == habito.id and r.fecha == fecha), None)
        if existente is not None:
            if not reemplazar:
                raise ErrorValidacion(f"Ya existe un registro de «{habito.nombre}» para {ftexto_fecha(fecha)}.")
            existente.cumplido, existente.nota = bool(cumplido), (nota or "").strip()
            registro = existente
        else:
            registro = Registro(fecha, habito.id, cumplido, nota)
            self._registros.append(registro)
        self.recalcular_rachas()
        self._persistir()
        return registro

    def eliminar_registro(self, ref: Union[int, str], fecha: date) -> bool:
        habito = self.buscar_habito(ref)
        antes = len(self._registros)
        self._registros = [r for r in self._registros
                           if not (r.id_habito == habito.id and r.fecha == fecha)]
        eliminado = len(self._registros) < antes
        if eliminado:
            self.recalcular_rachas()
            self._persistir()
        return eliminado

    def estado_dia(self, ref: Union[int, str], fecha: date) -> Optional[bool]:
        """True = cumplido, False = no cumplido, None = sin registro."""
        habito = self.buscar_habito(ref)
        for r in self._registros:
            if r.id_habito == habito.id and r.fecha == fecha:
                return r.cumplido
        return None

    def alternar_dia(self, ref: Union[int, str], fecha: date) -> Optional[bool]:
        """Ciclo de un clic en el calendario: sin registro -> Sí -> No -> sin registro."""
        actual = self.estado_dia(ref, fecha)
        if actual is None:
            self.registrar(ref, True, fecha)
            return True
        if actual:
            self.registrar(ref, False, fecha, reemplazar=True)
            return False
        self.eliminar_registro(ref, fecha)
        return None

    def recalcular_rachas(self) -> None:
        """RachaCorrida por FECHA (no por posición), como la fórmula del Excel:
        0 si el día no se cumplió; si se cumplió, 1 + racha del registro anterior
        de ese mismo hábito (el de la fecha inmediatamente menor)."""
        anterior: Dict[int, int] = {}
        for r in sorted(self._registros, key=lambda x: x.fecha):   # orden estable
            r.racha = (anterior.get(r.id_habito, 0) + 1) if r.cumplido else 0
            anterior[r.id_habito] = r.racha

    # ---------------------------------------------------------- reportes y modelo
    def reporte(self) -> Reporte:
        self.recalcular_rachas()
        return Reporte(self._habitos, self._registros, self.hoy())

    def modelo_vs_real(self, ref: Union[int, str]) -> ModeloVsReal:
        self.recalcular_rachas()
        return ModeloVsReal(self.modelo, self.buscar_habito(ref), self._registros)

    # ---------------------------------------------------------- persistencia
    def a_dict(self) -> dict:
        return {"version": 1, "usuario": self.usuario.a_dict(), "modelo": self.modelo.a_dict(),
                "habitos": [h.a_dict() for h in self._habitos],
                "registros": [r.a_dict() for r in self._registros]}

    def desde_dict(self, datos: dict) -> None:
        self.usuario = Usuario.desde_dict(datos.get("usuario", {}))
        m = datos.get("modelo", {})
        self.modelo = ModeloAutomaticidad(m.get("a_max", 100), m.get("dias_95", 66), m.get("epsilon", 1))
        self._habitos = [Habito.desde_dict(h) for h in datos.get("habitos", [])]
        self._registros = [Registro.desde_dict(r) for r in datos.get("registros", [])]
        self.recalcular_rachas()

    def cargar(self) -> bool:
        datos = self.repositorio.cargar()
        if datos is None:
            return False
        self.desde_dict(datos)
        return True

    def guardar(self) -> None:
        self.repositorio.guardar(self.a_dict())

    def _persistir(self) -> None:
        if self.autoguardar:
            self.guardar()

    def cargar_ejemplo(self) -> None:
        """Reemplaza los datos por los del Excel de ejemplo."""
        self.usuario = Usuario("Estudiante CUN", "Consolidar 3 hábitos clave este semestre", date(2026, 9, 28))
        self.modelo = ModeloAutomaticidad()
        self._habitos = [Habito(i, n, c, Frecuencia.DIARIA, _EJEMPLO_INICIO)
                         for i, (n, c, _) in enumerate(_EJEMPLO_HABITOS, start=1)]
        self._registros = []
        for dia in range(30):
            for i, (nombre, _, patron) in enumerate(_EJEMPLO_HABITOS, start=1):
                nota = _EJEMPLO_NOTAS.get((nombre, dia + 1), "")
                self._registros.append(Registro(_EJEMPLO_INICIO + timedelta(days=dia), i,
                                                patron[dia] == "S", nota))
        self.recalcular_rachas()
        self._persistir()

    def exportar_csv(self, ruta: Union[str, Path]) -> Path:
        self.recalcular_rachas()
        ruta = Path(ruta)
        nombres = {h.id: h.nombre for h in self._habitos}
        with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Fecha", "Hábito", "Cumplido", "Nota", "RachaCorrida"])
            for r in sorted(self._registros, key=lambda x: (x.fecha, x.id_habito)):
                w.writerow([r.fecha.isoformat(), nombres.get(r.id_habito, "?"),
                            "Sí" if r.cumplido else "No", r.nota, r.racha])
        return ruta


# ======================================================================
# 9. Puente con el Excel original (opcional: requiere openpyxl)
# ======================================================================


class PuenteExcel:
    """Lee los datos del libro .xlsx y compara sus resultados con los de Python."""

    @staticmethod
    def _abrir(ruta: Union[str, Path]):
        try:
            import openpyxl
        except ImportError:
            raise ErrorTracker("Para leer el Excel instala openpyxl:  pip install openpyxl")
        if not Path(ruta).exists():
            raise ErrorTracker(f"No existe el archivo: {ruta}")
        return openpyxl.load_workbook(str(ruta), data_only=True)

    @staticmethod
    def _como_fecha(valor) -> Optional[date]:
        if isinstance(valor, datetime):
            return valor.date()
        return valor if isinstance(valor, date) else None

    @classmethod
    def importar(cls, tracker: Tracker, ruta: Union[str, Path]) -> None:
        """Copia usuario, parámetros del modelo, hábitos y registros del Excel al Tracker."""
        wb = cls._abrir(ruta)
        cfg, cal = wb["Configuración"], wb["Cálculo"]
        tracker.usuario = Usuario(cfg["B3"].value or "Estudiante", cfg["B4"].value or "",
                                  cls._como_fecha(cfg["B5"].value))
        tracker.modelo = ModeloAutomaticidad(cal["B9"].value or 100, cal["B10"].value or 66,
                                             cal["B12"].value or 1)
        tracker._habitos, tracker._registros = [], []
        gh, rd = wb["Gestión de Hábitos"], wb["Registro Diario"]
        fila = 4
        while gh.cell(fila, 2).value:
            id_h, nombre, cat, frec, creado, estado = (gh.cell(fila, c).value for c in range(1, 7))
            try:
                habito = Habito(id_h, nombre, cat or "Otro", Frecuencia.desde_texto(frec or "Diaria"),
                                cls._como_fecha(creado), Estado.desde_texto(estado or "Activo"))
            except ErrorValidacion:
                habito = Habito(id_h, nombre, "Otro", Frecuencia.DIARIA, cls._como_fecha(creado))
            tracker._habitos.append(habito)
            fila += 1
        por_nombre = {h.nombre: h.id for h in tracker._habitos}
        fila = 4
        while rd.cell(fila, 1).value is not None:
            fecha, nombre, cumplido, nota = (rd.cell(fila, c).value for c in range(1, 5))
            if nombre in por_nombre and cls._como_fecha(fecha):
                tracker._registros.append(Registro(cls._como_fecha(fecha), por_nombre[nombre],
                                                   str(cumplido).strip().lower() in ("sí", "si"), nota or ""))
            fila += 1
        tracker.recalcular_rachas()

    @classmethod
    def comparar(cls, tracker: Tracker, ruta: Union[str, Path]) -> Tuple[str, bool]:
        """Devuelve (texto de comparación, ¿todo coincide?). Usa los valores que el Excel
        tiene guardados (se guardan al abrir y grabar el libro en Excel)."""
        wb = cls._abrir(ruta)
        est, dash, mvr = wb["Estadísticas"], wb["Dashboard"], wb["Modelo vs Real"]
        hoy_excel = cls._como_fecha(dash["B4"].value)
        if est["B4"].value is None:
            return ("El Excel no tiene valores guardados (nunca se calculó). Ábrelo en Excel, "
                    "guárdalo y vuelve a ejecutar la comparación.", False)
        if hoy_excel:
            tracker.fijar_hoy(hoy_excel)       # TODAY() del Excel, para que los 7/30 días coincidan
        reporte = tracker.reporte()
        filas, todo_ok = [], True

        def agregar(concepto: str, valor_excel, valor_py, decimales: int = 6) -> None:
            nonlocal todo_ok
            if isinstance(valor_excel, datetime):
                valor_excel = valor_excel.date()
            ok = (valor_excel == valor_py) if not isinstance(valor_py, float) or valor_excel is None \
                else abs(float(valor_excel) - valor_py) < 10 ** (-decimales)
            todo_ok = todo_ok and ok
            filas.append((concepto, str(valor_excel), str(valor_py), "OK" if ok else "DIFERENTE"))

        fila = 4
        while est.cell(fila, 1).value:
            habito = next((h for h in tracker.habitos if h.nombre == est.cell(fila, 1).value), None)
            if habito:
                e = reporte.estadistica(habito)
                nombre = habito.nombre
                agregar(f"{nombre} · racha actual", est.cell(fila, 2).value, e.racha_actual)
                agregar(f"{nombre} · mejor racha", est.cell(fila, 3).value, e.mejor_racha)
                agregar(f"{nombre} · % total", est.cell(fila, 4).value, e.pct_total)
                agregar(f"{nombre} · % 7 días", est.cell(fila, 5).value, e.pct_7_dias)
                agregar(f"{nombre} · % 30 días", est.cell(fila, 6).value, e.pct_30_dias)
            fila += 1
        agregar("Mejor racha global", est["B8"].value, reporte.mejor_racha_global(False))
        agregar("Promedio de cumplimiento", est["B9"].value, reporte.promedio_cumplimiento(False))
        agregar("k", wb["Cálculo"]["B11"].value, tracker.modelo.k)
        agregar("t*", wb["Cálculo"]["B13"].value, tracker.modelo.t_estrella)
        nombre_sel = mvr["B4"].value
        try:
            res = tracker.modelo_vs_real(nombre_sel).resultado()
        except ErrorValidacion:
            res = None
        if res:
            agregar(f"Modelo vs Real ({nombre_sel}) · T", mvr["B8"].value, res.dias_transcurridos)
            agregar("   n (repeticiones)", mvr["B9"].value, res.repeticiones)
            agregar("   ritmo n/T", mvr["B10"].value, res.ritmo)
            agregar("   A(n)", mvr["B11"].value, res.automaticidad_hoy)
            agregar("   repeticiones faltantes", mvr["B12"].value, float(res.repeticiones_faltantes))
            agregar("   días restantes", mvr["B13"].value, res.dias_restantes)
            agregar("   fecha de consolidación", mvr["B14"].value, res.fecha_consolidacion)
        texto = f"Comparación con {Path(ruta).name} (HOY del Excel: {ftexto_fecha(hoy_excel)})\n\n"
        texto += tabla(("Concepto", "Excel", "Python", "Resultado"), filas)
        texto += "\n\n" + ("✔ Todos los valores coinciden." if todo_ok else "✘ Hay diferencias: revisa las filas marcadas.")
        return texto, todo_ok


# ======================================================================
# 10. Gráficos (opcional: requiere matplotlib)
# ======================================================================


class Graficador:
    """Genera los mismos gráficos del libro de Excel como imágenes PNG."""

    def __init__(self, tracker: Tracker):
        self.tracker = tracker

    def generar(self, carpeta: Union[str, Path]) -> List[Path]:
        try:
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt
        except ImportError:
            raise ErrorTracker("Para los gráficos instala matplotlib:  pip install matplotlib")
        carpeta = Path(carpeta)
        carpeta.mkdir(parents=True, exist_ok=True)
        verde, ambar, oscuro = "#1E7B4F", "#E0A526", "#12261C"
        archivos: List[Path] = []
        est = self.tracker.reporte().estadisticas()
        nombres = [e.habito.nombre for e in est]

        fig, ax = plt.subplots(figsize=(8, 4.5))
        ax.bar(nombres, [e.racha_actual for e in est], color=verde)
        ax.set_title("Dashboard: racha actual por hábito")
        ax.set_ylabel("días seguidos")
        plt.setp(ax.get_xticklabels(), rotation=15, ha="right")
        fig.tight_layout()
        archivos.append(self._guardar(fig, carpeta / "dashboard_rachas.png"))

        fig, ax = plt.subplots(figsize=(8, 4.5))
        ax.bar(nombres, [e.pct_total * 100 for e in est], color=ambar)
        ax.set_ylim(0, 100)
        ax.set_title("Dashboard: % de cumplimiento total")
        ax.set_ylabel("%")
        plt.setp(ax.get_xticklabels(), rotation=15, ha="right")
        fig.tight_layout()
        archivos.append(self._guardar(fig, carpeta / "dashboard_porcentaje.png"))

        m = self.tracker.modelo
        filas = m.tabla(91)
        fig, ax = plt.subplots(figsize=(8, 4.5))
        ax.plot([f[0] for f in filas], [f[1] for f in filas], color=verde, label="A(t)")
        ax.plot([f[0] for f in filas], [f[2] for f in filas], color=ambar, label="A'(t) (puntos/día)")
        ax.axhline(m.a_max, color=oscuro, linestyle="--", label="A_máx")
        ax.axvline(m.t_estrella, color="gray", linestyle=":", label=f"t* ≈ {fnum(m.t_estrella)} días")
        ax.set_title("Cálculo: modelo de automaticidad")
        ax.set_xlabel("día (t)")
        ax.legend()
        fig.tight_layout()
        archivos.append(self._guardar(fig, carpeta / "curva_modelo.png"))

        activos = self.tracker.habitos_activos()
        if activos and self.tracker.registros_de(activos[0]):
            serie = self.tracker.modelo_vs_real(activos[0].id).serie(90)
            fig, ax = plt.subplots(figsize=(8, 4.5))
            ax.plot([f["t"] for f in serie], [f["real"] for f in serie], color=verde,
                    label=f"{activos[0].nombre} (real y proyección)")
            ax.plot([f["t"] for f in serie], [f["ideal"] for f in serie], color=ambar,
                    linestyle="--", label="Ideal (cumples todos los días)")
            ax.axhline(m.a_max, color=oscuro, linestyle=":", label="A_máx")
            ax.set_title("Modelo vs Real")
            ax.set_xlabel("día (t)")
            ax.set_ylabel("automaticidad (%)")
            ax.legend()
            fig.tight_layout()
            archivos.append(self._guardar(fig, carpeta / "modelo_vs_real.png"))
        return archivos

    @staticmethod
    def _guardar(fig, ruta: Path) -> Path:
        import matplotlib.pyplot as plt
        fig.savefig(ruta, dpi=120)
        plt.close(fig)
        return ruta


# ======================================================================
# 11. Interfaz de consola (la "vista"): separada de la lógica
# ======================================================================

AYUDA = """\
MANUAL DE USO — Tracker de Hábitos Inteligente (versión Python)

1. Agregar un hábito: menú «Gestión de hábitos» > Nuevo. Elige categoría y frecuencia.
2. Registrar el día: «Registrar hoy» pregunta por cada hábito activo que aún no tenga
   registro de hoy. «Registrar otra fecha» sirve para días pasados.
   La racha (RachaCorrida) se calcula sola, por fecha: no importa el orden en que registres.
3. Indicadores:
   • Racha actual: racha del registro más reciente del hábito (si el último día fue «No», es 0).
   • Mejor racha: la racha más larga que ha alcanzado.
   • % de cumplimiento: días cumplidos ÷ días registrados (total, últimos 7 y 30 días).
4. Dashboard: resumen de todos los hábitos activos.
5. Cálculo: modelo A(t) = A_máx·(1 − e^(−k·t)), su derivada A'(t), su segunda derivada A''(t)
   y el día t* en que la ganancia diaria baja de ε.
6. Modelo vs Real: lleva las repeticiones cumplidas (n) de un hábito al modelo y estima
   cuántos días faltan para el 95 %.
7. Los datos se guardan solos en tracker_datos.json (junto a este programa). Haz copias de
   seguridad de ese archivo.
"""


class SalirPrograma(Exception):
    """Se lanza cuando el usuario cierra la entrada (Ctrl+D / Ctrl+Z)."""


class MenuConsola:
    def __init__(self, tracker: Tracker):
        self.t = tracker

    # ------------------------------------------------------ entrada de datos
    @staticmethod
    def leer(mensaje: str) -> str:
        try:
            return input(mensaje)
        except EOFError:
            raise SalirPrograma()

    def pedir_texto(self, mensaje: str, defecto: Optional[str] = None) -> str:
        sufijo = f" [{defecto}]" if defecto else ""
        valor = self.leer(f"{mensaje}{sufijo}: ").strip()
        return valor or (defecto or "")

    def pedir_fecha(self, mensaje: str, defecto: Optional[date] = None) -> date:
        valor = self.leer(f"{mensaje} [{(defecto or self.t.hoy()).isoformat()}]: ").strip()
        return parsear_fecha(valor) if valor else (defecto or self.t.hoy())

    def pedir_numero(self, mensaje: str, defecto: float) -> float:
        valor = self.leer(f"{mensaje} [{fnum(defecto, 4).rstrip('0').rstrip(',')}]: ").strip()
        if not valor:
            return defecto
        try:
            return float(valor.replace(",", "."))
        except ValueError:
            raise ErrorValidacion(f"«{valor}» no es un número.")

    def elegir(self, titulo: str, opciones: Sequence[str]) -> int:
        print(titulo)
        for i, o in enumerate(opciones, start=1):
            print(f"  {i}. {o}")
        valor = self.leer("Elige un número: ").strip()
        if not valor.isdigit() or not (1 <= int(valor) <= len(opciones)):
            raise ErrorValidacion("Opción no válida.")
        return int(valor) - 1

    def elegir_habito(self, solo_activos: bool = False) -> Habito:
        lista = self.t.habitos_activos() if solo_activos else list(self.t.habitos)
        if not lista:
            raise ErrorValidacion("Todavía no hay hábitos. Crea uno en «Gestión de hábitos».")
        return lista[self.elegir("Hábitos:", [f"{h.nombre} ({h.estado.value})" for h in lista])]

    # ------------------------------------------------------ bucle principal
    def ejecutar(self) -> None:
        opciones = [
            ("Dashboard", self.dashboard),
            ("Registrar hoy", self.registrar_hoy),
            ("Registrar otra fecha", self.registrar_fecha),
            ("Gestión de hábitos", self.gestion_habitos),
            ("Ver registro diario", self.ver_registros),
            ("Estadísticas por hábito", self.estadisticas),
            ("Cálculo: modelo de automaticidad", self.calculo),
            ("Modelo vs Real", self.modelo_vs_real),
            ("Configuración", self.configuracion),
            ("Exportar (CSV y gráficos)", self.exportar),
            ("Manual / Ayuda", lambda: print(AYUDA)),
        ]
        print("\n" + "=" * 62)
        print("  TRACKER DE HÁBITOS INTELIGENTE  ·  versión Python (POO)")
        print("=" * 62)
        try:
            while True:
                print(f"\nHoy: {ftexto_fecha(self.t.hoy())}   Usuario: {self.t.usuario.nombre}")
                for i, (titulo, _) in enumerate(opciones, start=1):
                    print(f"  {i:>2}. {titulo}")
                print("   0. Salir")
                eleccion = self.leer("\nElige una opción: ").strip()
                if eleccion == "0":
                    break
                if not eleccion.isdigit() or not (1 <= int(eleccion) <= len(opciones)):
                    print("Opción no válida.")
                    continue
                try:
                    print()
                    opciones[int(eleccion) - 1][1]()
                except ErrorTracker as error:
                    print(f"⚠ {error}")
        except (SalirPrograma, KeyboardInterrupt):
            print()
        self.t.guardar()
        print("Datos guardados. ¡Hasta pronto!")

    # ------------------------------------------------------ pantallas
    def dashboard(self) -> None:
        print(texto_dashboard(self.t))

    def registrar_hoy(self) -> None:
        hoy = self.t.hoy()
        activos = self.t.habitos_activos()
        if not activos:
            raise ErrorValidacion("No hay hábitos activos.")
        pendientes = [h for h in activos
                      if not any(r.fecha == hoy for r in self.t.registros_de(h))]
        if not pendientes:
            print("Todos los hábitos activos ya tienen registro de hoy.")
            return
        registrados = 0
        for h in pendientes:
            resp = self.leer(f"¿Cumpliste hoy «{h.nombre}»? [s]í / [n]o / Enter=omitir / [c]ancelar: ").strip().lower()
            if resp in ("c", "cancelar"):
                print("Registro cancelado. Lo ya guardado se conserva.")
                break
            if resp in ("s", "si", "sí", "n", "no"):
                self.t.registrar(h.id, resp in ("s", "si", "sí"), hoy)
                registrados += 1
        print(f"{registrados} hábito(s) registrado(s) para hoy ({ftexto_fecha(hoy)}).")

    def registrar_fecha(self) -> None:
        h = self.elegir_habito(solo_activos=True)
        fecha = self.pedir_fecha("Fecha")
        cumplido = self.leer("¿Cumplido? [s/n]: ").strip().lower() in ("s", "si", "sí")
        nota = self.pedir_texto("Nota (opcional)")
        try:
            self.t.registrar(h.id, cumplido, fecha, nota)
        except ErrorValidacion as error:
            if "Ya existe un registro" in str(error) and \
                    self.leer("Ya existe. ¿Reemplazarlo? [s/n]: ").strip().lower() in ("s", "si", "sí"):
                self.t.registrar(h.id, cumplido, fecha, nota, reemplazar=True)
            else:
                raise
        print("Registro guardado.")

    def gestion_habitos(self) -> None:
        accion = self.elegir("Gestión de hábitos", ["Ver lista", "Nuevo hábito", "Editar hábito",
                                                    "Archivar hábito", "Reactivar hábito"])
        if accion == 0:
            print(texto_habitos(self.t))
        elif accion == 1:
            nombre = self.pedir_texto("Nombre del hábito")
            cat = CATEGORIAS[self.elegir("Categoría:", CATEGORIAS)]
            frec = list(Frecuencia)[self.elegir("Frecuencia:", [f.value for f in Frecuencia])]
            h = self.t.crear_habito(nombre, cat, frec)
            print(f"Hábito «{h.nombre}» creado con id {h.id}.")
        elif accion == 2:
            h = self.elegir_habito()
            nombre = self.pedir_texto("Nuevo nombre", h.nombre)
            cat = self.pedir_texto(f"Categoría ({', '.join(CATEGORIAS)})", h.categoria)
            frec = Frecuencia.desde_texto(self.pedir_texto("Frecuencia (Diaria/Semanal/Mensual)", h.frecuencia.value))
            self.t.editar_habito(h.id, nombre, cat, frec)
            print("Hábito actualizado.")
        elif accion == 3:
            h = self.elegir_habito()
            self.t.archivar_habito(h.id)
            print(f"«{h.nombre}» archivado (su historial se conserva).")
        else:
            h = self.elegir_habito()
            self.t.reactivar_habito(h.id)
            print(f"«{h.nombre}» reactivado.")

    def ver_registros(self) -> None:
        opciones = ["Todos los hábitos"] + [h.nombre for h in self.t.habitos]
        eleccion = self.elegir("¿De qué hábito?", opciones)
        cuantos = int(self.pedir_numero("¿Cuántos registros recientes mostrar?", 15))
        print(texto_registros(self.t, None if eleccion == 0 else self.t.habitos[eleccion - 1], cuantos))

    def estadisticas(self) -> None:
        print(texto_estadisticas(self.t))

    def calculo(self) -> None:
        print(texto_calculo(self.t.modelo))
        if self.leer("\n¿Editar los parámetros del modelo? [s/n]: ").strip().lower() in ("s", "si", "sí"):
            m = self.t.modelo
            m.configurar(self.pedir_numero("A_máx (nivel máximo, %)", m.a_max),
                         self.pedir_numero("Días para alcanzar el 95 %", m.dias_95),
                         self.pedir_numero("ε: ganancia diaria mínima (puntos/día)", m.epsilon))
            self.t.guardar()
            print("\n" + texto_calculo(m))

    def modelo_vs_real(self) -> None:
        h = self.elegir_habito()
        print(texto_modelo_vs_real(self.t.modelo_vs_real(h.id)))

    def configuracion(self) -> None:
        u = self.t.usuario
        print(f"Usuario: {u.nombre}\nMeta general: {u.meta_general}\n"
              f"Inicio del sistema: {ftexto_fecha(u.fecha_inicio)}\n")
        u.configurar(self.pedir_texto("Nombre del usuario", u.nombre),
                     self.pedir_texto("Meta general", u.meta_general))
        self.t.guardar()
        print("Configuración guardada.")

    def exportar(self) -> None:
        carpeta = Path(__file__).resolve().parent
        csv_ruta = self.t.exportar_csv(carpeta / "registros_exportados.csv")
        print(f"CSV guardado en: {csv_ruta}")
        try:
            archivos = Graficador(self.t).generar(carpeta / "graficos")
            print("Gráficos guardados en:\n  " + "\n  ".join(str(a) for a in archivos))
        except ErrorTracker as error:
            print(f"(Gráficos omitidos) {error}")


# ------------------------------------------------------------------
# Textos de salida (funciones puras: devuelven texto, no imprimen)
# ------------------------------------------------------------------


def texto_dashboard(t: Tracker) -> str:
    rep = t.reporte()
    est = rep.estadisticas()
    lineas = ["=== DASHBOARD — Tracker de Hábitos Inteligente ===",
              f"Usuario: {t.usuario.nombre}   Meta: {t.usuario.meta_general}",
              f"Fecha de hoy: {ftexto_fecha(t.hoy())}",
              f"Mejor racha global: {rep.mejor_racha_global()} días   "
              f"Promedio de cumplimiento: {fpct(rep.promedio_cumplimiento())}", ""]
    if not est:
        return "\n".join(lineas + ["(sin hábitos activos)"])
    filas = [(e.habito.nombre, e.racha_actual, barra(e.pct_total), fpct(e.pct_total)) for e in est]
    lineas.append(tabla(("Hábito", "Racha", "Cumplimiento", "%"), filas, ("<", ">", "<", ">")))
    return "\n".join(lineas)


def texto_habitos(t: Tracker) -> str:
    filas = [(h.id, h.nombre, h.categoria, h.frecuencia.value, ftexto_fecha(h.fecha_creacion), h.estado.value)
             for h in t.habitos]
    return tabla(("ID", "Nombre", "Categoría", "Frecuencia", "Creación", "Estado"), filas) if filas \
        else "(no hay hábitos)"


def texto_registros(t: Tracker, habito: Optional[Habito], cuantos: int) -> str:
    nombres = {h.id: h.nombre for h in t.habitos}
    t.recalcular_rachas()
    regs = [r for r in t.registros if habito is None or r.id_habito == habito.id]
    regs = sorted(regs, key=lambda r: (r.fecha, r.id_habito))[-max(1, cuantos):]
    filas = [(ftexto_fecha(r.fecha), nombres.get(r.id_habito, "?"), "Sí" if r.cumplido else "No",
              r.racha, r.nota) for r in regs]
    return tabla(("Fecha", "Hábito", "Cumplido", "Racha", "Nota"), filas, ("<", "<", "<", ">", "<")) \
        if filas else "(sin registros)"


def texto_estadisticas(t: Tracker) -> str:
    rep = t.reporte()
    filas = [(e.habito.nombre, e.racha_actual, e.mejor_racha, fpct(e.pct_total),
              fpct(e.pct_7_dias), fpct(e.pct_30_dias)) for e in rep.estadisticas()]
    cuerpo = tabla(("Hábito", "Racha actual", "Mejor racha", "% total", "% 7 días", "% 30 días"),
                   filas, ("<", ">", ">", ">", ">", ">")) if filas else "(sin hábitos activos)"
    return (f"=== ESTADÍSTICAS POR HÁBITO ===  (hoy: {ftexto_fecha(t.hoy())})\n{cuerpo}\n\n"
            f"Mejor racha global: {rep.mejor_racha_global()}\n"
            f"Promedio % cumplimiento: {fpct(rep.promedio_cumplimiento())}")


def texto_calculo(m: ModeloAutomaticidad) -> str:
    lineas = [
        "=== CÁLCULO: modelo de automaticidad de hábitos ===",
        "Modelo:  A(t)   = A_máx · (1 − e^(−k·t))        (viene de  A'(t) = k·(A_máx − A),  A(0) = 0)",
        "         A'(t)  = A_máx · k · e^(−k·t)           (puntos ganados por día)",
        "         A''(t) = −A_máx · k² · e^(−k·t) = −k·A'(t)   (negativa: el avance frena)",
        "",
        f"A_máx (nivel máximo, %)                  = {fnum(m.a_max, 2)}",
        f"Días para alcanzar el 95 % (Lally 2010)  = {fnum(m.dias_95, 2)}",
        f"k = ln(20) / {fnum(m.dias_95, 0)}                         = {fnum(m.k, 6)}",
        f"ε (ganancia diaria mínima)               = {fnum(m.epsilon, 2)}",
        f"t* = ln(A_máx·k/ε) / k                   = {fnum(m.t_estrella, 2)} días"
        f"   (A(t*) = {fnum(m.automaticidad(m.t_estrella), 2)} %)",
        "",
    ]
    dias = sorted({0, 5, 10, 15, 20, 30, 40, 50, 66, 80, 91, int(round(m.dias_95))})
    filas = [(t, fnum(m.automaticidad(t), 2), fnum(m.velocidad(t), 3), fnum(m.aceleracion(t), 4),
              barra(m.automaticidad(t) / m.a_max, 30)) for t in dias]
    lineas.append(tabla(("Día t", "A(t)", "A'(t)", "A''(t)", "Curva"), filas, (">", ">", ">", ">", "<")))
    return "\n".join(lineas)


def texto_modelo_vs_real(mvr: ModeloVsReal) -> str:
    res = mvr.resultado()
    if res is None:
        return f"«{mvr.habito.nombre}» todavía no tiene registros."
    lineas = [
        f"=== MODELO vs REAL — {res.habito.nombre} ===",
        f"Primer registro:                      {ftexto_fecha(res.primer_registro)}",
        f"Último registro:                      {ftexto_fecha(res.ultimo_registro)}",
        f"Días transcurridos (T):               {res.dias_transcurridos}",
        f"Repeticiones cumplidas (n):           {res.repeticiones}",
        f"Ritmo de cumplimiento (n / T):        {fnum(res.ritmo, 3)}",
        f"Automaticidad estimada hoy, A(n):     {fnum(res.automaticidad_hoy, 2)} %",
        f"Repeticiones que faltan para el 95 %: {fnum(res.repeticiones_faltantes, 0)}",
        f"Días restantes estimados (tu ritmo):  {res.dias_restantes if res.dias_restantes is not None else 'sin datos'}",
        f"Fecha estimada de consolidación:      {ftexto_fecha(res.fecha_consolidacion)}",
        "", "Día a día (real hasta el último registro, luego proyección):"]
    serie = mvr.serie(90)
    muestra = [f for f in serie if f["t"] in (1, 5, 10, 15, 20, 25, 30, 40, 50, 60, 70, 80, 90)]
    filas = [(f["t"], ftexto_fecha(f["fecha"]),
              "" if f["pct_real"] is None else fpct(f["pct_real"], 0),
              fnum(f["n"], 1), fnum(f["real"], 1), fnum(f["ideal"], 1),
              "proyección" if f["proyeccion"] else "real") for f in muestra]
    lineas.append(tabla(("Día", "Fecha", "% acum.", "n(t)", "Tu hábito", "Ideal", ""), filas,
                        (">", "<", ">", ">", ">", ">", "<")))
    return "\n".join(lineas)


# ======================================================================
# 11b. Interfaz gráfica (Tkinter): ventana con pestañas, clics y gráficas
# ======================================================================
# Tkinter viene incluido con Python (en Windows y macOS no hay que instalar nada).
# La ventana usa las MISMAS clases de arriba (Tracker, Reporte, Modelo…): solo cambia
# la "vista". Las gráficas se dibujan a mano sobre un Canvas, sin librerías externas.

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    HAY_TK = True
except ImportError:                      # Python sin Tk (algunos Linux): se usa la consola
    HAY_TK = False

VERDE, AMBAR, OSCURO = "#1E7B4F", "#E0A526", "#12261C"
FONDO, BLANCO, BORDE, GRIS = "#F5F4EE", "#FFFFFF", "#DAD8CC", "#6A7A70"
ROJO, ROJO_CLARO, VACIO = "#C0463A", "#E8A29A", "#E6E4DA"
FUENTE = "Segoe UI"
DIAS_SEMANA = ("L", "M", "X", "J", "V", "S", "D")
NOMBRES_DIA = ("Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom")
MESES = ("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic")

AYUDA_GUI = """\
MANUAL — Tracker de Hábitos Inteligente (versión con interfaz)

PANEL (Dashboard)
  • Arriba: indicadores generales. En «Registrar hoy» pulsa «Hecho» o «No» en cada hábito.
  • A la derecha: gráficas de racha actual y de % de cumplimiento (se actualizan solas).

CALENDARIO
  • Elige un hábito. Cada cuadro es un día: verde = cumplido, rojo = no cumplido, gris = sin registro.
  • Haz CLIC en un día para cambiarlo: sin registro → cumplido → no cumplido → sin registro.
  • Abajo puedes escribir una nota, cambiar la fecha o eliminar un registro.
  • La racha se calcula por fecha: no importa el orden en que registres.

HÁBITOS
  • Selecciona uno para editarlo, o escribe uno nuevo y pulsa «Nuevo hábito».
  • Archivar un hábito conserva todo su historial; reactivarlo lo devuelve al panel.

ESTADÍSTICAS
  • Racha actual, mejor racha y % de cumplimiento (total, últimos 7 y últimos 30 días).

CÁLCULO (Cálculo Diferencial)
  • Mueve los deslizadores: la curva A(t) = A_máx·(1 − e^(−k·t)) y su derivada A'(t) cambian en vivo.
  • Pasa el mouse sobre la gráfica para leer A, A' y A'' de cualquier día.
  • k = ln(20) / días para el 95 %.  t* es el día en que ganas menos de ε puntos por día.

MODELO vs REAL
  • Lleva tus repeticiones reales (n) a la curva y estima cuántos días faltan para el 95 %.
  • Activa «Simular escenario» y mueve el deslizador: ¿qué pasa si desde hoy cumplo el X % de los días?

CONFIGURACIÓN
  • Cambia tu nombre y meta, carga los datos de ejemplo, exporta CSV / gráficos PNG,
    o importa y compara con el libro de Excel original.

Los datos se guardan solos en tracker_datos.json (junto a este programa). Haz copias de ese archivo.
"""


def _pasos(maximo: float) -> List[float]:
    """Marcas «bonitas» para un eje que va de 0 a `maximo`."""
    if maximo <= 0:
        return [0.0]
    bruto = maximo / 5
    magnitud = 10 ** math.floor(math.log10(bruto))
    paso = next(m * magnitud for m in (1, 2, 2.5, 5, 10) if m * magnitud >= bruto)
    return [i * paso for i in range(int(maximo / paso + 1e-9) + 1)]


class Plano:
    """Convierte coordenadas de datos a píxeles y dibuja ejes sobre un Canvas."""

    def __init__(self, c, w: int, h: int, xr: Tuple[float, float], yr: Tuple[float, float],
                 yr2: Optional[Tuple[float, float]] = None,
                 izq: int = 54, der: int = 16, arr: int = 16, aba: int = 36):
        self.c, self.w, self.h = c, w, h
        self.x0, self.x1 = xr
        self.y0, self.y1 = yr
        self.z0, self.z1 = yr2 if yr2 else (0.0, 1.0)
        self.izq, self.arr = izq, arr
        self.der = 56 if yr2 else der
        self.aba = aba
        self.ancho = max(10, w - self.izq - self.der)
        self.alto = max(10, h - self.arr - self.aba)

    def px(self, x: float) -> float:
        return self.izq + (x - self.x0) / (self.x1 - self.x0) * self.ancho

    def py(self, y: float) -> float:
        return self.arr + (1 - (y - self.y0) / (self.y1 - self.y0)) * self.alto

    def py2(self, y: float) -> float:
        return self.arr + (1 - (y - self.z0) / (self.z1 - self.z0)) * self.alto

    def x_de_pixel(self, p: float) -> float:
        return self.x0 + (p - self.izq) / self.ancho * (self.x1 - self.x0)

    def dentro(self, p: float) -> bool:
        return self.izq <= p <= self.izq + self.ancho

    def marco(self, xt: Sequence[float], yt: Sequence[float], xlabel: str = "", ylabel: str = "",
              yt2: Optional[Sequence[float]] = None, ylabel2: str = "") -> None:
        c = self.c
        for v in yt:
            y = self.py(v)
            c.create_line(self.izq, y, self.izq + self.ancho, y, fill=BORDE)
            c.create_text(self.izq - 6, y, text=f"{v:g}", anchor="e", fill=GRIS, font=(FUENTE, 8))
        for v in xt:
            x = self.px(v)
            c.create_line(x, self.arr + self.alto, x, self.arr + self.alto + 4, fill=GRIS)
            c.create_text(x, self.arr + self.alto + 14, text=f"{v:g}", fill=GRIS, font=(FUENTE, 8))
        if yt2:
            for v in yt2:
                c.create_text(self.izq + self.ancho + 6, self.py2(v), text=f"{v:g}", anchor="w",
                              fill=AMBAR, font=(FUENTE, 8))
        c.create_line(self.izq, self.arr, self.izq, self.arr + self.alto, fill=GRIS)
        c.create_line(self.izq, self.arr + self.alto, self.izq + self.ancho, self.arr + self.alto, fill=GRIS)
        if xlabel:
            c.create_text(self.izq + self.ancho / 2, self.h - 8, text=xlabel, fill=OSCURO, font=(FUENTE, 9))
        if ylabel:
            c.create_text(4, self.arr - 2, text=ylabel, anchor="nw", fill=VERDE, font=(FUENTE, 9, "bold"))
        if ylabel2:
            c.create_text(self.w - 4, self.arr - 2, text=ylabel2, anchor="ne", fill="#9A7010",
                          font=(FUENTE, 9, "bold"))

    def curva(self, xs: Sequence[float], ys: Sequence[float], color: str, ancho: int = 3,
              dash: Optional[Tuple[int, int]] = None, eje2: bool = False) -> None:
        if len(xs) < 2:
            return
        f = self.py2 if eje2 else self.py
        pts: List[float] = []
        for x, y in zip(xs, ys):
            pts += [self.px(x), f(y)]
        self.c.create_line(*pts, fill=color, width=ancho, dash=dash, capstyle="round", joinstyle="round")

    def vertical(self, x: float, color: str, dash=(4, 3), texto: str = "", arriba: bool = True) -> None:
        p = self.px(x)
        self.c.create_line(p, self.arr, p, self.arr + self.alto, fill=color, dash=dash)
        if texto:
            self.c.create_text(p + 4, self.arr + (4 if arriba else self.alto - 12), text=texto,
                               anchor="nw", fill=color, font=(FUENTE, 8, "bold"))

    def horizontal(self, y: float, color: str, dash=(4, 3), texto: str = "") -> None:
        p = self.py(y)
        self.c.create_line(self.izq, p, self.izq + self.ancho, p, fill=color, dash=dash)
        if texto:
            self.c.create_text(self.izq + self.ancho - 4, p - 2, text=texto, anchor="se",
                               fill=color, font=(FUENTE, 8, "bold"))


def dibujar_barras(c, w: int, h: int, titulo: str, etiquetas: Sequence[str],
                   series: Sequence[Tuple[str, Sequence[float], str]], maximo: float, sufijo: str = "") -> None:
    """Barras agrupadas con valores encima. series = [(nombre, valores, color), ...]."""
    c.create_text(10, 8, text=titulo, anchor="nw", fill=OSCURO, font=(FUENTE, 11, "bold"))
    if not etiquetas:
        c.create_text(w / 2, h / 2, text="(sin hábitos activos)", fill=GRIS, font=(FUENTE, 10))
        return
    p = Plano(c, w, h, (0, len(etiquetas)), (0, maximo), izq=40, der=12, arr=34, aba=34)
    p.marco([], _pasos(maximo))
    grupo = 0.72 / len(series)
    for i, etiqueta in enumerate(etiquetas):
        c.create_text(p.px(i + 0.5), p.arr + p.alto + 16,
                      text=etiqueta if len(etiqueta) <= 16 else etiqueta[:15] + "…",
                      fill=OSCURO, font=(FUENTE, 8))
        for j, (_, valores, color) in enumerate(series):
            x0 = p.px(i + 0.14 + j * grupo)
            x1 = p.px(i + 0.14 + (j + 1) * grupo) - 2
            y = p.py(valores[i])
            c.create_rectangle(x0, y, x1, p.py(0), fill=color, outline="")
            c.create_text((x0 + x1) / 2, y - 7, text=f"{valores[i]:g}{sufijo}", fill=OSCURO, font=(FUENTE, 8))
    if len(series) > 1:
        x = w - 12
        for nombre, _, color in reversed(series):
            c.create_text(x, 10, text=nombre, anchor="ne", fill=OSCURO, font=(FUENTE, 8))
            x -= len(nombre) * 6 + 22
            c.create_rectangle(x, 6, x + 10, 16, fill=color, outline="")
            x -= 8


class Lienzo(tk.Canvas if HAY_TK else object):          # type: ignore[misc]
    """Canvas que se vuelve a dibujar solo cuando cambia de tamaño o se pide refrescar."""

    def __init__(self, master, dibujar: Callable, **kw):
        super().__init__(master, bg=BLANCO, highlightthickness=1, highlightbackground=BORDE, **kw)
        self._dibujar = dibujar
        self.bind("<Configure>", lambda _e: self.redibujar())

    def redibujar(self) -> None:
        self.delete("all")
        w, h = self.winfo_width(), self.winfo_height()
        if w > 60 and h > 60:
            self._dibujar(self, w, h)


class AplicacionGrafica(tk.Tk if HAY_TK else object):    # type: ignore[misc]
    """Ventana principal. Una pestaña por hoja del Excel; todo se refresca con refrescar()."""

    SEMANAS_CALENDARIO = 20

    def __init__(self, tracker: Tracker, primera_vez: bool = False):
        super().__init__()
        self.t = tracker
        self.title("Tracker de Hábitos Inteligente — Python (POO)")
        self.geometry("1200x780")
        self.minsize(1000, 660)
        self.configure(bg=FONDO)
        self._celdas: List[Tuple[float, float, float, float, date]] = []
        self._plano_calculo: Optional[Plano] = None
        self._plano_mvr: Optional[Plano] = None
        self._guia = None
        self._estilos()
        self._cabecera()
        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        self._pestana_panel()
        self._pestana_calendario()
        self._pestana_habitos()
        self._pestana_estadisticas()
        self._pestana_calculo()
        self._pestana_modelo_real()
        self._pestana_configuracion()
        self._pestana_ayuda()
        self.protocol("WM_DELETE_WINDOW", self._cerrar)
        self.refrescar()
        if primera_vez:
            self.after(400, self._bienvenida)

    # ------------------------------------------------------------ estilos y cabecera
    def _estilos(self) -> None:
        s = ttk.Style(self)
        try:
            s.theme_use("clam")
        except tk.TclError:
            pass
        s.configure(".", font=(FUENTE, 10), background=FONDO)
        s.configure("TFrame", background=FONDO)
        s.configure("Tarjeta.TFrame", background=BLANCO, relief="solid", borderwidth=1)
        s.configure("TLabel", background=FONDO, foreground=OSCURO)
        s.configure("Tarjeta.TLabel", background=BLANCO)
        s.configure("TNotebook", background=FONDO, borderwidth=0)
        s.configure("TNotebook.Tab", padding=(16, 8), font=(FUENTE, 10, "bold"),
                    background=BORDE, foreground=OSCURO)
        s.map("TNotebook.Tab", background=[("selected", BLANCO)], foreground=[("selected", VERDE)])
        s.configure("Treeview", rowheight=26, font=(FUENTE, 10), background=BLANCO, fieldbackground=BLANCO)
        s.configure("Treeview.Heading", font=(FUENTE, 10, "bold"), background=OSCURO, foreground=FONDO)
        s.map("Treeview", background=[("selected", VERDE)], foreground=[("selected", BLANCO)])
        s.configure("TButton", padding=(12, 6))
        s.configure("Acento.TButton", background=VERDE, foreground=BLANCO, font=(FUENTE, 10, "bold"))
        s.map("Acento.TButton", background=[("active", "#17613E")])
        s.configure("Peligro.TButton", background=ROJO, foreground=BLANCO, font=(FUENTE, 10, "bold"))
        s.map("Peligro.TButton", background=[("active", "#9C382E")])
        s.configure("TLabelframe", background=FONDO)
        s.configure("TLabelframe.Label", background=FONDO, foreground=VERDE, font=(FUENTE, 10, "bold"))
        s.configure("TCheckbutton", background=FONDO)

    def _cabecera(self) -> None:
        barra_ = tk.Frame(self, bg=OSCURO)
        barra_.pack(fill="x")
        tk.Label(barra_, text="Tracker de Hábitos Inteligente", bg=OSCURO, fg=FONDO,
                 font=(FUENTE, 17, "bold")).pack(side="left", padx=16, pady=10)
        self.lbl_cab = tk.Label(barra_, text="", bg=OSCURO, fg=AMBAR, font=(FUENTE, 10, "bold"))
        self.lbl_cab.pack(side="left", padx=10)
        ttk.Button(barra_, text="Guardar", style="Acento.TButton",
                   command=self._seguro(self._guardar)).pack(side="right", padx=16, pady=8)

    # ------------------------------------------------------------ utilidades
    def _seguro(self, funcion: Callable) -> Callable:
        """Envuelve un callback: los errores del dominio salen como aviso, no como traza."""
        def envuelta(*args):
            try:
                funcion(*args)
            except ErrorTracker as error:
                messagebox.showwarning("Aviso", str(error), parent=self)
        return envuelta

    def _guardar(self) -> None:
        self.t.guardar()
        self.lbl_cab.config(text=self._texto_cabecera() + "   ·   Guardado ✔")
        self.after(1800, lambda: self.lbl_cab.config(text=self._texto_cabecera()))

    def _cerrar(self) -> None:
        try:
            self.t.guardar()
        finally:
            self.destroy()

    def _texto_cabecera(self) -> str:
        return f"Hola, {self.t.usuario.nombre}   ·   Hoy: {ftexto_fecha(self.t.hoy())}"

    def _bienvenida(self) -> None:
        if messagebox.askyesno("Bienvenido", "Aún no hay datos guardados.\n\n¿Cargar los datos de ejemplo "
                               "del Excel (3 hábitos × 30 días)?\n\n«No» empieza con todo vacío.", parent=self):
            self.t.cargar_ejemplo()
        else:
            self.t.guardar()
        self.refrescar()

    def _nombres(self, solo_activos: bool) -> List[str]:
        return [h.nombre for h in (self.t.habitos_activos() if solo_activos else self.t.habitos)]

    @staticmethod
    def _ajustar_combo(combo: "ttk.Combobox", valores: List[str]) -> None:
        actual = combo.get()
        combo["values"] = valores
        combo.set(actual if actual in valores else (valores[0] if valores else ""))

    # ------------------------------------------------------------ pestaña: panel
    def _pestana_panel(self) -> None:
        f = ttk.Frame(self.nb, padding=12)
        self.nb.add(f, text="  Panel  ")
        fila = ttk.Frame(f)
        fila.pack(fill="x")
        self.kpis: Dict[str, tk.Label] = {}
        for i, (clave, titulo) in enumerate((("activos", "Hábitos activos"), ("hoy", "Cumplidos hoy"),
                                             ("mejor", "Mejor racha global"), ("prom", "Cumplimiento promedio"))):
            tarjeta = ttk.Frame(fila, style="Tarjeta.TFrame", padding=(16, 10))
            tarjeta.grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else 10, 0))
            fila.columnconfigure(i, weight=1)
            ttk.Label(tarjeta, text=titulo, style="Tarjeta.TLabel", foreground=GRIS).pack(anchor="w")
            self.kpis[clave] = tk.Label(tarjeta, text="—", bg=BLANCO, fg=VERDE, font=(FUENTE, 24, "bold"))
            self.kpis[clave].pack(anchor="w")

        cuerpo = ttk.Frame(f)
        cuerpo.pack(fill="both", expand=True, pady=(12, 0))
        cuerpo.columnconfigure(0, weight=5, uniform="c")
        cuerpo.columnconfigure(1, weight=6, uniform="c")
        cuerpo.rowconfigure(0, weight=1)
        izq = ttk.LabelFrame(cuerpo, text="Registrar hoy", padding=10)
        izq.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        self.f_hoy = ttk.Frame(izq)
        self.f_hoy.pack(fill="both", expand=True)
        der = ttk.Frame(cuerpo)
        der.grid(row=0, column=1, sticky="nsew")
        der.rowconfigure((0, 1), weight=1)
        der.columnconfigure(0, weight=1)
        self.g_rachas = Lienzo(der, self._dib_rachas)
        self.g_rachas.grid(row=0, column=0, sticky="nsew", pady=(0, 10))
        self.g_pct = Lienzo(der, self._dib_porcentaje)
        self.g_pct.grid(row=1, column=0, sticky="nsew")

    def _dib_rachas(self, c, w, h) -> None:
        est = self.t.reporte().estadisticas()
        mayor = max([e.mejor_racha for e in est] + [5])
        dibujar_barras(c, w, h, "Racha actual y mejor racha (días seguidos)", [e.habito.nombre for e in est],
                       [("Actual", [e.racha_actual for e in est], VERDE),
                        ("Mejor", [e.mejor_racha for e in est], AMBAR)], math.ceil(mayor * 1.15))

    def _dib_porcentaje(self, c, w, h) -> None:
        est = self.t.reporte().estadisticas()
        dibujar_barras(c, w, h, "% de cumplimiento", [e.habito.nombre for e in est],
                       [("Total", [round(e.pct_total * 100) for e in est], VERDE),
                        ("30 días", [round(e.pct_30_dias * 100) for e in est], AMBAR),
                        ("7 días", [round(e.pct_7_dias * 100) for e in est], "#8FB9A3")], 110, "")

    def _refrescar_panel(self) -> None:
        rep = self.t.reporte()
        est = rep.estadisticas()
        hoy = self.t.hoy()
        hechos = sum(1 for e in est if self.t.estado_dia(e.habito.id, hoy) is True)
        self.kpis["activos"].config(text=str(len(est)))
        self.kpis["hoy"].config(text=f"{hechos} / {len(est)}")
        self.kpis["mejor"].config(text=f"{rep.mejor_racha_global()} días")
        self.kpis["prom"].config(text=fpct(rep.promedio_cumplimiento(), 0))
        for hijo in self.f_hoy.winfo_children():
            hijo.destroy()
        if not est:
            ttk.Label(self.f_hoy, text="No hay hábitos activos.\nCrea uno en la pestaña «Hábitos».",
                      foreground=GRIS).pack(pady=30)
        for i, e in enumerate(est):
            estado = self.t.estado_dia(e.habito.id, hoy)
            tarjeta = ttk.Frame(self.f_hoy, style="Tarjeta.TFrame", padding=(10, 8))
            tarjeta.pack(fill="x", pady=4)
            tarjeta.columnconfigure(0, weight=1)
            ttk.Label(tarjeta, text=e.habito.nombre, style="Tarjeta.TLabel",
                      font=(FUENTE, 11, "bold")).grid(row=0, column=0, sticky="w")
            texto, color = (("Hoy: ✔ cumplido", VERDE) if estado is True else
                            ("Hoy: ✘ no cumplido", ROJO) if estado is False else ("Hoy: pendiente", GRIS))
            tk.Label(tarjeta, text=texto, bg=BLANCO, fg=color, font=(FUENTE, 9, "bold")).grid(
                row=1, column=0, sticky="w")
            tk.Label(tarjeta, text=f"Racha {e.racha_actual}  ·  {fpct(e.pct_total, 0)}", bg=BLANCO, fg=GRIS,
                     font=(FUENTE, 9)).grid(row=2, column=0, sticky="w")
            ttk.Button(tarjeta, text="Hecho", style="Acento.TButton",
                       command=self._seguro(lambda h=e.habito: self._marcar_hoy(h, True))).grid(
                row=0, column=1, rowspan=3, padx=(6, 4))
            ttk.Button(tarjeta, text="No", style="Peligro.TButton",
                       command=self._seguro(lambda h=e.habito: self._marcar_hoy(h, False))).grid(
                row=0, column=2, rowspan=3)
        self.g_rachas.redibujar()
        self.g_pct.redibujar()

    def _marcar_hoy(self, habito: Habito, cumplido: bool) -> None:
        self.t.registrar(habito.id, cumplido, self.t.hoy(), reemplazar=True)
        self.refrescar()

    # ------------------------------------------------------------ pestaña: calendario / registro
    def _pestana_calendario(self) -> None:
        f = ttk.Frame(self.nb, padding=12)
        self.nb.add(f, text="  Calendario  ")
        sup = ttk.Frame(f)
        sup.pack(fill="x")
        ttk.Label(sup, text="Hábito:", font=(FUENTE, 10, "bold")).pack(side="left")
        self.cb_cal = ttk.Combobox(sup, state="readonly", width=30)
        self.cb_cal.pack(side="left", padx=8)
        self.cb_cal.bind("<<ComboboxSelected>>", lambda _e: self._refrescar_calendario())
        ttk.Label(sup, text="Clic en un día:  sin registro > cumplido > no cumplido > sin registro", foreground=GRIS).pack(
            side="left", padx=16)
        self.lbl_tip = ttk.Label(sup, text="", foreground=VERDE, font=(FUENTE, 10, "bold"))
        self.lbl_tip.pack(side="right")

        self.cal = Lienzo(f, self._dib_calendario, height=250)
        self.cal.pack(fill="x", pady=10)
        self.cal.bind("<Button-1>", self._seguro(self._clic_calendario))
        self.cal.bind("<Motion>", self._mover_calendario)
        self.cal.bind("<Leave>", lambda _e: self.lbl_tip.config(text=""))

        inf = ttk.Frame(f)
        inf.pack(fill="both", expand=True)
        inf.columnconfigure(0, weight=3)
        inf.columnconfigure(1, weight=2)
        inf.rowconfigure(0, weight=1)
        marco_t = ttk.Frame(inf)
        marco_t.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        self.tv_reg = ttk.Treeview(marco_t, columns=("fecha", "cumplido", "racha", "nota"), show="headings",
                                   selectmode="browse")
        for col, texto, ancho in (("fecha", "Fecha", 100), ("cumplido", "Cumplido", 90),
                                  ("racha", "Racha", 70), ("nota", "Nota", 260)):
            self.tv_reg.heading(col, text=texto)
            self.tv_reg.column(col, width=ancho, anchor="w" if col == "nota" else "center")
        sb = ttk.Scrollbar(marco_t, orient="vertical", command=self.tv_reg.yview)
        self.tv_reg.configure(yscrollcommand=sb.set)
        self.tv_reg.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.tv_reg.tag_configure("si", foreground=VERDE)
        self.tv_reg.tag_configure("no", foreground=ROJO)
        self.tv_reg.bind("<<TreeviewSelect>>", self._seleccion_registro)

        form = ttk.LabelFrame(inf, text="Registrar / editar un día", padding=10)
        form.grid(row=0, column=1, sticky="nsew")
        form.columnconfigure(1, weight=1)
        ttk.Label(form, text="Fecha").grid(row=0, column=0, sticky="w", pady=4)
        fr = ttk.Frame(form)
        fr.grid(row=0, column=1, sticky="ew")
        self.var_fecha = tk.StringVar()
        ttk.Button(fr, text="<", width=3, command=self._seguro(lambda: self._mover_fecha(-1))).pack(side="left")
        ttk.Entry(fr, textvariable=self.var_fecha, width=12, justify="center").pack(side="left", padx=4)
        ttk.Button(fr, text=">", width=3, command=self._seguro(lambda: self._mover_fecha(1))).pack(side="left")
        ttk.Button(fr, text="Hoy", width=5,
                   command=lambda: self.var_fecha.set(self.t.hoy().isoformat())).pack(side="left", padx=6)
        self.var_cumplido = tk.BooleanVar(value=True)
        ttk.Checkbutton(form, text="Cumplido", variable=self.var_cumplido).grid(row=1, column=1, sticky="w", pady=4)
        ttk.Label(form, text="Nota").grid(row=2, column=0, sticky="w", pady=4)
        self.var_nota = tk.StringVar()
        ttk.Entry(form, textvariable=self.var_nota).grid(row=2, column=1, sticky="ew")
        ttk.Button(form, text="Guardar registro", style="Acento.TButton",
                   command=self._seguro(self._guardar_registro)).grid(row=3, column=0, columnspan=2,
                                                                       sticky="ew", pady=(12, 4))
        ttk.Button(form, text="Eliminar este día", style="Peligro.TButton",
                   command=self._seguro(self._eliminar_registro)).grid(row=4, column=0, columnspan=2, sticky="ew")
        self.var_fecha.set(self.t.hoy().isoformat())

    def _habito_calendario(self) -> Optional[Habito]:
        nombre = self.cb_cal.get()
        return next((h for h in self.t.habitos if h.nombre == nombre), None)

    def _dib_calendario(self, c, w, h) -> None:
        habito = self._habito_calendario()
        self._celdas = []
        if habito is None:
            c.create_text(w / 2, h / 2, text="Crea un hábito para ver su calendario", fill=GRIS, font=(FUENTE, 11))
            return
        hoy = self.t.hoy()
        inicio = hoy - timedelta(days=hoy.weekday()) - timedelta(weeks=self.SEMANAS_CALENDARIO - 1)
        lado = max(14, min(34, (w - 70) / self.SEMANAS_CALENDARIO, (h - 44) / 7))
        x_base, y_base = 40, 28
        estados = {r.fecha: r for r in self.t.registros_de(habito)}
        for j, letra in enumerate(DIAS_SEMANA):
            c.create_text(x_base - 14, y_base + j * lado + lado / 2, text=letra, fill=GRIS, font=(FUENTE, 8))
        mes_previo = -1
        for semana in range(self.SEMANAS_CALENDARIO):
            lunes = inicio + timedelta(weeks=semana)
            if lunes.month != mes_previo:
                c.create_text(x_base + semana * lado, 12, text=MESES[lunes.month - 1], anchor="w",
                              fill=OSCURO, font=(FUENTE, 8, "bold"))
                mes_previo = lunes.month
            for j in range(7):
                dia = lunes + timedelta(days=j)
                x0, y0 = x_base + semana * lado, y_base + j * lado
                x1, y1 = x0 + lado - 3, y0 + lado - 3
                if dia > hoy:
                    c.create_rectangle(x0, y0, x1, y1, fill=FONDO, outline=BORDE, dash=(2, 2))
                    continue
                r = estados.get(dia)
                color = VACIO if r is None else (VERDE if r.cumplido else ROJO_CLARO)
                c.create_rectangle(x0, y0, x1, y1, fill=color, outline=OSCURO if dia == hoy else "",
                                   width=2 if dia == hoy else 1)
                if lado >= 22:
                    c.create_text((x0 + x1) / 2, (y0 + y1) / 2, text=str(dia.day),
                                  fill=BLANCO if (r and r.cumplido) else GRIS, font=(FUENTE, 8))
                self._celdas.append((x0, y0, x1, y1, dia))
        lx = x_base + self.SEMANAS_CALENDARIO * lado + 12
        for k, (texto, color) in enumerate((("Cumplido", VERDE), ("No cumplido", ROJO_CLARO), ("Sin registro", VACIO))):
            c.create_rectangle(lx, y_base + k * 22, lx + 12, y_base + k * 22 + 12, fill=color, outline="")
            c.create_text(lx + 18, y_base + k * 22 + 6, text=texto, anchor="w", fill=OSCURO, font=(FUENTE, 8))

    def _celda_en(self, x: float, y: float) -> Optional[date]:
        for x0, y0, x1, y1, dia in self._celdas:
            if x0 <= x <= x1 and y0 <= y <= y1:
                return dia
        return None

    def _clic_calendario(self, evento) -> None:
        dia, habito = self._celda_en(evento.x, evento.y), self._habito_calendario()
        if dia and habito:
            self.t.alternar_dia(habito.id, dia)
            self.var_fecha.set(dia.isoformat())
            self.refrescar()

    def _mover_calendario(self, evento) -> None:
        dia, habito = self._celda_en(evento.x, evento.y), self._habito_calendario()
        if not (dia and habito):
            self.lbl_tip.config(text="")
            return
        reg = next((r for r in self.t.registros_de(habito) if r.fecha == dia), None)
        estado = "sin registro" if reg is None else ("cumplido · racha " + str(reg.racha) if reg.cumplido
                                                      else "no cumplido")
        nota = f" · «{reg.nota}»" if reg and reg.nota else ""
        self.lbl_tip.config(text=f"{NOMBRES_DIA[dia.weekday()]} {ftexto_fecha(dia)} — {estado}{nota}")

    def _refrescar_calendario(self) -> None:
        habito = self._habito_calendario()
        self.cal.redibujar()
        self.tv_reg.delete(*self.tv_reg.get_children())
        if habito:
            for r in reversed(self.t.registros_de(habito)):
                self.tv_reg.insert("", "end", iid=r.fecha.isoformat(),
                                   values=(ftexto_fecha(r.fecha), "Sí" if r.cumplido else "No", r.racha, r.nota),
                                   tags=("si" if r.cumplido else "no",))

    def _seleccion_registro(self, _e=None) -> None:
        sel = self.tv_reg.selection()
        habito = self._habito_calendario()
        if sel and habito:
            fecha = date.fromisoformat(sel[0])
            reg = next(r for r in self.t.registros_de(habito) if r.fecha == fecha)
            self.var_fecha.set(fecha.isoformat())
            self.var_cumplido.set(reg.cumplido)
            self.var_nota.set(reg.nota)

    def _mover_fecha(self, delta: int) -> None:
        self.var_fecha.set((parsear_fecha(self.var_fecha.get()) + timedelta(days=delta)).isoformat())

    def _guardar_registro(self) -> None:
        habito = self._habito_calendario()
        if habito is None:
            raise ErrorValidacion("Primero crea o elige un hábito.")
        self.t.registrar(habito.id, self.var_cumplido.get(), parsear_fecha(self.var_fecha.get()),
                         self.var_nota.get(), reemplazar=True)
        self.refrescar()

    def _eliminar_registro(self) -> None:
        habito = self._habito_calendario()
        if habito is None:
            raise ErrorValidacion("Primero elige un hábito.")
        fecha = parsear_fecha(self.var_fecha.get())
        if not self.t.eliminar_registro(habito.id, fecha):
            raise ErrorValidacion(f"No hay registro de {ftexto_fecha(fecha)} para eliminar.")
        self.refrescar()

    # ------------------------------------------------------------ pestaña: hábitos
    def _pestana_habitos(self) -> None:
        f = ttk.Frame(self.nb, padding=12)
        self.nb.add(f, text="  Hábitos  ")
        f.columnconfigure(0, weight=3)
        f.columnconfigure(1, weight=2)
        f.rowconfigure(0, weight=1)
        self.tv_hab = ttk.Treeview(f, columns=("id", "nombre", "cat", "frec", "creado", "estado"),
                                   show="headings", selectmode="browse")
        for col, texto, ancho in (("id", "ID", 40), ("nombre", "Nombre", 220), ("cat", "Categoría", 110),
                                  ("frec", "Frecuencia", 90), ("creado", "Creación", 100), ("estado", "Estado", 90)):
            self.tv_hab.heading(col, text=texto)
            self.tv_hab.column(col, width=ancho, anchor="w" if col == "nombre" else "center")
        self.tv_hab.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        self.tv_hab.tag_configure("archivado", foreground=GRIS)
        self.tv_hab.bind("<<TreeviewSelect>>", self._seleccion_habito)

        form = ttk.LabelFrame(f, text="Hábito", padding=12)
        form.grid(row=0, column=1, sticky="new")
        form.columnconfigure(1, weight=1)
        self.var_h_nombre = tk.StringVar()
        self.var_h_cat = tk.StringVar(value=CATEGORIAS[0])
        self.var_h_frec = tk.StringVar(value=Frecuencia.DIARIA.value)
        ttk.Label(form, text="Nombre").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Entry(form, textvariable=self.var_h_nombre).grid(row=0, column=1, sticky="ew")
        ttk.Label(form, text="Categoría").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Combobox(form, textvariable=self.var_h_cat, values=CATEGORIAS, state="readonly").grid(
            row=1, column=1, sticky="ew")
        ttk.Label(form, text="Frecuencia").grid(row=2, column=0, sticky="w", pady=4)
        ttk.Combobox(form, textvariable=self.var_h_frec, values=[x.value for x in Frecuencia],
                     state="readonly").grid(row=2, column=1, sticky="ew")
        ttk.Button(form, text="Nuevo hábito", style="Acento.TButton",
                   command=self._seguro(self._nuevo_habito)).grid(row=3, column=0, columnspan=2, sticky="ew", pady=(14, 4))
        ttk.Button(form, text="Guardar cambios del seleccionado",
                   command=self._seguro(self._editar_habito)).grid(row=4, column=0, columnspan=2, sticky="ew", pady=4)
        ttk.Button(form, text="Archivar / Reactivar seleccionado",
                   command=self._seguro(self._archivar_habito)).grid(row=5, column=0, columnspan=2, sticky="ew", pady=4)
        ttk.Label(form, text="Archivar conserva el historial.\nRenombrar no pierde registros.",
                  foreground=GRIS).grid(row=6, column=0, columnspan=2, sticky="w", pady=(10, 0))

    def _habito_seleccionado(self) -> Habito:
        sel = self.tv_hab.selection()
        if not sel:
            raise ErrorValidacion("Selecciona un hábito de la lista.")
        return self.t.buscar_habito(int(sel[0]))

    def _seleccion_habito(self, _e=None) -> None:
        sel = self.tv_hab.selection()
        if sel:
            h = self.t.buscar_habito(int(sel[0]))
            self.var_h_nombre.set(h.nombre)
            self.var_h_cat.set(h.categoria)
            self.var_h_frec.set(h.frecuencia.value)

    def _nuevo_habito(self) -> None:
        h = self.t.crear_habito(self.var_h_nombre.get(), self.var_h_cat.get(),
                                Frecuencia.desde_texto(self.var_h_frec.get()))
        self.var_h_nombre.set("")
        self.refrescar()
        self.cb_cal.set(h.nombre)
        self._refrescar_calendario()

    def _editar_habito(self) -> None:
        h = self._habito_seleccionado()
        self.t.editar_habito(h.id, self.var_h_nombre.get(), self.var_h_cat.get(),
                             Frecuencia.desde_texto(self.var_h_frec.get()))
        self.refrescar()

    def _archivar_habito(self) -> None:
        h = self._habito_seleccionado()
        (self.t.archivar_habito if h.esta_activo else self.t.reactivar_habito)(h.id)
        self.refrescar()

    def _refrescar_habitos(self) -> None:
        seleccionado = self.tv_hab.selection()
        self.tv_hab.delete(*self.tv_hab.get_children())
        for h in self.t.habitos:
            self.tv_hab.insert("", "end", iid=str(h.id), tags=("" if h.esta_activo else "archivado",),
                               values=(h.id, h.nombre, h.categoria, h.frecuencia.value,
                                       ftexto_fecha(h.fecha_creacion), h.estado.value))
        if seleccionado and self.tv_hab.exists(seleccionado[0]):
            self.tv_hab.selection_set(seleccionado[0])

    # ------------------------------------------------------------ pestaña: estadísticas
    def _pestana_estadisticas(self) -> None:
        f = ttk.Frame(self.nb, padding=12)
        self.nb.add(f, text="  Estadísticas  ")
        self.tv_est = ttk.Treeview(f, columns=("h", "ra", "mr", "pt", "p7", "p30", "reg"), show="headings",
                                   height=6, selectmode="none")
        for col, texto, ancho in (("h", "Hábito", 240), ("ra", "Racha actual", 100), ("mr", "Mejor racha", 100),
                                  ("pt", "% total", 90), ("p7", "% 7 días", 90), ("p30", "% 30 días", 90),
                                  ("reg", "Cumplidos / registros", 150)):
            self.tv_est.heading(col, text=texto)
            self.tv_est.column(col, width=ancho, anchor="w" if col == "h" else "center")
        self.tv_est.pack(fill="x")
        self.lbl_est = ttk.Label(f, text="", font=(FUENTE, 11, "bold"), foreground=VERDE)
        self.lbl_est.pack(anchor="w", pady=10)
        self.g_est = Lienzo(f, self._dib_estadisticas)
        self.g_est.pack(fill="both", expand=True)

    def _dib_estadisticas(self, c, w, h) -> None:
        est = self.t.reporte().estadisticas()
        dibujar_barras(c, w, h, "Cumplimiento por periodo (%)", [e.habito.nombre for e in est],
                       [("Total", [round(e.pct_total * 100, 1) for e in est], VERDE),
                        ("Últimos 30 días", [round(e.pct_30_dias * 100, 1) for e in est], AMBAR),
                        ("Últimos 7 días", [round(e.pct_7_dias * 100, 1) for e in est], "#8FB9A3")], 110)

    def _refrescar_estadisticas(self) -> None:
        rep = self.t.reporte()
        self.tv_est.delete(*self.tv_est.get_children())
        for e in rep.estadisticas():
            self.tv_est.insert("", "end", values=(e.habito.nombre, e.racha_actual, e.mejor_racha,
                                                  fpct(e.pct_total), fpct(e.pct_7_dias), fpct(e.pct_30_dias),
                                                  f"{e.cumplidos} / {e.registros}"))
        self.lbl_est.config(text=f"Mejor racha global: {rep.mejor_racha_global()} días     "
                                 f"Promedio de cumplimiento: {fpct(rep.promedio_cumplimiento())}")
        self.g_est.redibujar()

    # ------------------------------------------------------------ pestaña: cálculo
    def _pestana_calculo(self) -> None:
        f = ttk.Frame(self.nb, padding=12)
        self.nb.add(f, text="  Cálculo  ")
        f.columnconfigure(1, weight=1)
        f.rowconfigure(0, weight=1)
        panel = ttk.LabelFrame(f, text="Parámetros del modelo", padding=12)
        panel.grid(row=0, column=0, sticky="ns", padx=(0, 12))
        m = self.t.modelo
        self.esc_amax = self._deslizador(panel, "A_máx (nivel máximo, %)", 10, 100, 1, m.a_max, 0)
        self.esc_dias = self._deslizador(panel, "Días para llegar al 95 %", 10, 200, 1, m.dias_95, 1)
        self.esc_eps = self._deslizador(panel, "ε (ganancia mínima por día)", 0.1, 5, 0.1, m.epsilon, 2)
        ttk.Button(panel, text="Valores del proyecto (100 · 66 · 1)", command=self._seguro(self._reset_modelo)).grid(
            row=6, column=0, sticky="ew", pady=10)
        self.lbl_formulas = tk.Label(panel, text="", bg=FONDO, fg=OSCURO, justify="left", anchor="w",
                                     font=("Consolas", 10))
        self.lbl_formulas.grid(row=7, column=0, sticky="w", pady=(6, 0))
        der = ttk.Frame(f)
        der.grid(row=0, column=1, sticky="nsew")
        der.rowconfigure(0, weight=1)
        der.columnconfigure(0, weight=1)
        self.g_calc = Lienzo(der, self._dib_calculo)
        self.g_calc.grid(row=0, column=0, sticky="nsew")
        self.g_calc.bind("<Motion>", lambda e: self._hover(e, self.g_calc, "calc"))
        self.g_calc.bind("<Leave>", lambda _e: self.lbl_calc.config(text=self._pista_calculo()))
        self.lbl_calc = ttk.Label(der, text=self._pista_calculo(), font=(FUENTE, 10, "bold"), foreground=VERDE)
        self.lbl_calc.grid(row=1, column=0, sticky="w", pady=(8, 0))

    @staticmethod
    def _pista_calculo() -> str:
        return "Pasa el mouse sobre la gráfica para leer A, A' y A'' de cada día."

    def _deslizador(self, padre, titulo: str, minimo: float, maximo: float, paso: float,
                    valor: float, fila: int) -> "tk.Scale":
        ttk.Label(padre, text=titulo, font=(FUENTE, 10, "bold")).grid(row=fila * 2, column=0, sticky="w", pady=(8, 0))
        escala = tk.Scale(padre, from_=minimo, to=maximo, resolution=paso, orient="horizontal", length=250,
                          bg=FONDO, highlightthickness=0, troughcolor=BORDE, activebackground=VERDE,
                          command=lambda _v: self._cambio_modelo())
        escala.set(valor)
        escala.grid(row=fila * 2 + 1, column=0, sticky="ew")
        escala.bind("<ButtonRelease-1>", lambda _e: self.t.guardar())
        return escala

    def _cambio_modelo(self) -> None:
        if not hasattr(self, "esc_eps"):
            return
        self.t.modelo.configurar(self.esc_amax.get(), self.esc_dias.get(), self.esc_eps.get())
        self._refrescar_modelo()

    def _reset_modelo(self) -> None:
        self.esc_amax.set(100)
        self.esc_dias.set(66)
        self.esc_eps.set(1)
        self.t.guardar()

    def _refrescar_modelo(self) -> None:
        m = self.t.modelo
        self.lbl_formulas.config(text=(
            f"A(t)  = A_máx·(1 − e^(−k·t))\nA'(t) = A_máx·k·e^(−k·t)\nA''(t)= −k·A'(t)\n\n"
            f"k  = ln(20)/{m.dias_95:g} = {fnum(m.k, 5)}\n"
            f"t* = ln(A_máx·k/ε)/k = {fnum(m.t_estrella, 1)} días\n"
            f"A(t*) = {fnum(m.automaticidad(m.t_estrella), 1)} %"))
        self.g_calc.redibujar()
        self.g_mvr.redibujar()
        self._refrescar_resultados_mvr()

    def _dib_calculo(self, c, w, h) -> None:
        m = self.t.modelo
        t_max = max(91, int(m.dias_95 * 1.4))
        ymax = max(10.0, m.a_max * 1.08)
        vmax = max(0.1, m.a_max * m.k * 1.1)
        p = Plano(c, w, h, (0, t_max), (0, ymax), (0, vmax))
        self._plano_calculo = p
        p.marco(_pasos(t_max), _pasos(ymax), "día (t)", "A(t)  — automaticidad (%)", _pasos(vmax), "A'(t)  — puntos/día")
        xs = list(range(0, t_max + 1))
        p.horizontal(m.a_max, OSCURO, texto=f"A_máx = {m.a_max:g}")
        p.horizontal(m.a_max * 0.95, GRIS, texto="95 %")
        p.vertical(m.dias_95, GRIS, texto=f"día {m.dias_95:g}", arriba=False)
        if m.t_estrella > 0:
            p.vertical(m.t_estrella, AMBAR, texto=f"t* = {fnum(m.t_estrella, 1)}")
        p.curva(xs, [m.automaticidad(t) for t in xs], VERDE, 3)
        p.curva(xs, [m.velocidad(t) for t in xs], AMBAR, 3, eje2=True)
        c.create_oval(p.px(m.dias_95) - 5, p.py(m.a_max * 0.95) - 5, p.px(m.dias_95) + 5, p.py(m.a_max * 0.95) + 5,
                      fill=VERDE, outline=BLANCO, width=2)
        self._guia = None

    # ------------------------------------------------------------ pestaña: modelo vs real
    def _pestana_modelo_real(self) -> None:
        f = ttk.Frame(self.nb, padding=12)
        self.nb.add(f, text="  Modelo vs Real  ")
        f.columnconfigure(1, weight=1)
        f.rowconfigure(0, weight=1)
        panel = ttk.LabelFrame(f, text="Hábito y escenario", padding=12)
        panel.grid(row=0, column=0, sticky="ns", padx=(0, 12))
        ttk.Label(panel, text="Hábito", font=(FUENTE, 10, "bold")).grid(row=0, column=0, sticky="w")
        self.cb_mvr = ttk.Combobox(panel, state="readonly", width=28)
        self.cb_mvr.grid(row=1, column=0, sticky="ew", pady=(2, 10))
        self.cb_mvr.bind("<<ComboboxSelected>>", lambda _e: self._refrescar_modelo())
        self.lbl_res = tk.Label(panel, text="", bg=FONDO, fg=OSCURO, justify="left", anchor="w", font=(FUENTE, 10))
        self.lbl_res.grid(row=2, column=0, sticky="w")
        self.var_sim = tk.BooleanVar(value=False)
        ttk.Checkbutton(panel, text="Simular escenario", variable=self.var_sim,
                        command=self._refrescar_modelo).grid(row=3, column=0, sticky="w", pady=(14, 0))
        self.esc_ritmo = tk.Scale(panel, from_=5, to=100, resolution=5, orient="horizontal", length=250,
                                  label="Desde hoy cumplo este % de los días", bg=FONDO, highlightthickness=0,
                                  troughcolor=BORDE, activebackground=VERDE,
                                  command=lambda _v: self._refrescar_modelo() if self.var_sim.get() else None)
        self.esc_ritmo.set(90)
        self.esc_ritmo.grid(row=4, column=0, sticky="ew")
        der = ttk.Frame(f)
        der.grid(row=0, column=1, sticky="nsew")
        der.rowconfigure(0, weight=1)
        der.columnconfigure(0, weight=1)
        self.g_mvr = Lienzo(der, self._dib_mvr)
        self.g_mvr.grid(row=0, column=0, sticky="nsew")
        self.g_mvr.bind("<Motion>", lambda e: self._hover(e, self.g_mvr, "mvr"))
        self.lbl_mvr = ttk.Label(der, text="Pasa el mouse sobre la gráfica para comparar día a día.",
                                 font=(FUENTE, 10, "bold"), foreground=VERDE)
        self.lbl_mvr.grid(row=1, column=0, sticky="w", pady=(8, 0))

    def _mvr(self) -> Optional[ModeloVsReal]:
        nombre = self.cb_mvr.get()
        return self.t.modelo_vs_real(nombre) if nombre else None

    def _ritmo_escenario(self) -> Optional[float]:
        return self.esc_ritmo.get() / 100 if self.var_sim.get() else None

    def _refrescar_resultados_mvr(self) -> None:
        mvr = self._mvr()
        res = mvr.resultado(self._ritmo_escenario()) if mvr else None
        if res is None:
            self.lbl_res.config(text="\nSin registros todavía:\nregistra algunos días\nen el Calendario.")
            return
        escenario = "  (escenario)" if self.var_sim.get() else ""
        self.lbl_res.config(text=(
            f"Primer registro:  {ftexto_fecha(res.primer_registro)}\n"
            f"Último registro:  {ftexto_fecha(res.ultimo_registro)}\n"
            f"Días transcurridos T:  {res.dias_transcurridos}\n"
            f"Repeticiones cumplidas n:  {res.repeticiones}\n"
            f"Ritmo{escenario}:  {fpct(res.ritmo, 0)}\n"
            f"Automaticidad hoy A(n):  {fnum(res.automaticidad_hoy, 1)} %\n"
            f"Faltan para el 95 %:  {fnum(res.repeticiones_faltantes, 0)} repeticiones\n"
            f"Días restantes:  {res.dias_restantes if res.dias_restantes is not None else 'sin datos'}\n"
            f"Consolidación:  {ftexto_fecha(res.fecha_consolidacion)}"))

    def _dib_mvr(self, c, w, h) -> None:
        mvr = self._mvr()
        m = self.t.modelo
        res = mvr.resultado(self._ritmo_escenario()) if mvr else None
        if res is None:
            c.create_text(w / 2, h / 2, text="Elige un hábito con registros", fill=GRIS, font=(FUENTE, 11))
            self._plano_mvr = None
            return
        ritmo = self._ritmo_escenario()
        serie = mvr.serie(90, ritmo)
        ymax = max(10.0, m.a_max * 1.08)
        p = Plano(c, w, h, (1, 90), (0, ymax))
        self._plano_mvr = p
        p.marco([1] + [x for x in range(10, 91, 10)], _pasos(ymax), "día del hábito (t)  ·  eje vertical: automaticidad (%)")
        p.horizontal(m.a_max, OSCURO, texto=f"A_máx = {m.a_max:g}")
        p.horizontal(m.a_max * 0.95, GRIS, texto="95 %")
        ts = [f["t"] for f in serie]
        p.curva(ts, [f["ideal"] for f in serie], AMBAR, 2, dash=(6, 4))
        reales = [f for f in serie if not f["proyeccion"]]
        proy = [f for f in serie if f["proyeccion"]]
        if reales:
            p.curva([f["t"] for f in reales], [f["real"] for f in reales], VERDE, 3)
        if proy:
            conecta = ([reales[-1]] if reales else []) + proy
            p.curva([f["t"] for f in conecta], [f["real"] for f in conecta], VERDE, 3, dash=(2, 4))
            p.vertical(res.dias_transcurridos, GRIS, texto="último registro")
        c.create_text(p.izq + 10, p.arr + 8, text="━ Tu hábito (real)", anchor="w", fill=VERDE, font=(FUENTE, 9, "bold"))
        c.create_text(p.izq + 140, p.arr + 8, text="┅ Proyección", anchor="w", fill=VERDE, font=(FUENTE, 9, "bold"))
        c.create_text(p.izq + 245, p.arr + 8, text="┅ Ideal (todos los días)", anchor="w", fill="#9A7010",
                      font=(FUENTE, 9, "bold"))
        if res.dias_restantes is not None and 0 < res.dias_transcurridos + res.dias_restantes <= 90:
            tx = res.dias_transcurridos + res.dias_restantes
            c.create_oval(p.px(tx) - 5, p.py(m.a_max * 0.95) - 5, p.px(tx) + 5, p.py(m.a_max * 0.95) + 5,
                          fill=VERDE, outline=BLANCO, width=2)
            c.create_text(p.px(tx) - 8, p.py(m.a_max * 0.95) + 14, text=f"95 % en el día {tx}", anchor="e",
                          fill=VERDE, font=(FUENTE, 9, "bold"))

    # ------------------------------------------------------------ hover común a las dos gráficas
    def _hover(self, evento, lienzo, cual: str) -> None:
        p = self._plano_calculo if cual == "calc" else self._plano_mvr
        lienzo.delete("guia")
        if p is None or not p.dentro(evento.x):
            return
        t = max(p.x0, min(p.x1, round(p.x_de_pixel(evento.x))))
        x = p.px(t)
        lienzo.create_line(x, p.arr, x, p.arr + p.alto, fill=OSCURO, dash=(2, 2), tags="guia")
        m = self.t.modelo
        if cual == "calc":
            lienzo.create_oval(x - 4, p.py(m.automaticidad(t)) - 4, x + 4, p.py(m.automaticidad(t)) + 4,
                               fill=VERDE, outline=BLANCO, tags="guia")
            lienzo.create_oval(x - 4, p.py2(m.velocidad(t)) - 4, x + 4, p.py2(m.velocidad(t)) + 4,
                               fill=AMBAR, outline=BLANCO, tags="guia")
            self.lbl_calc.config(text=f"Día {t}:   A = {fnum(m.automaticidad(t), 2)} %    "
                                      f"A' = {fnum(m.velocidad(t), 3)} pts/día    A'' = {fnum(m.aceleracion(t), 4)}")
        else:
            mvr = self._mvr()
            serie = mvr.serie(90, self._ritmo_escenario()) if mvr else []
            if serie and 1 <= t <= len(serie):
                f = serie[t - 1]
                lienzo.create_oval(x - 4, p.py(f["real"]) - 4, x + 4, p.py(f["real"]) + 4,
                                   fill=VERDE, outline=BLANCO, tags="guia")
                lienzo.create_oval(x - 4, p.py(f["ideal"]) - 4, x + 4, p.py(f["ideal"]) + 4,
                                   fill=AMBAR, outline=BLANCO, tags="guia")
                etiqueta = "proyección" if f["proyeccion"] else "real"
                self.lbl_mvr.config(text=f"Día {t} ({ftexto_fecha(f['fecha'])}, {etiqueta}):   n = {fnum(f['n'], 1)}   "
                                         f"tu hábito = {fnum(f['real'], 1)} %   ideal = {fnum(f['ideal'], 1)} %   "
                                         f"diferencia = {fnum(f['ideal'] - f['real'], 1)} pts")

    # ------------------------------------------------------------ pestaña: configuración
    def _pestana_configuracion(self) -> None:
        f = ttk.Frame(self.nb, padding=16)
        self.nb.add(f, text="  Configuración  ")
        f.columnconfigure(1, weight=1)
        self.var_nombre = tk.StringVar()
        self.var_meta = tk.StringVar()
        ttk.Label(f, text="Nombre del usuario", font=(FUENTE, 10, "bold")).grid(row=0, column=0, sticky="w", pady=6)
        ttk.Entry(f, textvariable=self.var_nombre, width=50).grid(row=0, column=1, sticky="w", padx=10)
        ttk.Label(f, text="Meta general", font=(FUENTE, 10, "bold")).grid(row=1, column=0, sticky="w", pady=6)
        ttk.Entry(f, textvariable=self.var_meta, width=50).grid(row=1, column=1, sticky="w", padx=10)
        self.lbl_inicio = ttk.Label(f, text="", foreground=GRIS)
        self.lbl_inicio.grid(row=2, column=1, sticky="w", padx=10)
        ttk.Button(f, text="Guardar configuración", style="Acento.TButton",
                   command=self._seguro(self._guardar_config)).grid(row=3, column=1, sticky="w", padx=10, pady=10)
        caja = ttk.LabelFrame(f, text="Datos", padding=12)
        caja.grid(row=4, column=0, columnspan=2, sticky="ew", pady=14)
        botones = (("Cargar datos de ejemplo", self._cargar_ejemplo),
                   ("Exportar registros a CSV", self._exportar_csv),
                   ("Generar gráficos PNG", self._exportar_graficos),
                   ("Importar desde el Excel…", self._importar_excel),
                   ("Comparar con el Excel…", self._comparar_excel))
        for i, (texto, accion) in enumerate(botones):
            ttk.Button(caja, text=texto, command=self._seguro(accion)).grid(row=0, column=i, padx=6, sticky="ew")
        ttk.Label(f, text="Importar y comparar necesitan openpyxl (pip install openpyxl); "
                          "los gráficos PNG necesitan matplotlib.", foreground=GRIS).grid(
            row=5, column=0, columnspan=2, sticky="w")

    def _guardar_config(self) -> None:
        self.t.usuario.configurar(self.var_nombre.get(), self.var_meta.get())
        self.t.guardar()
        self.refrescar()

    def _cargar_ejemplo(self) -> None:
        if messagebox.askyesno("Datos de ejemplo", "Esto REEMPLAZA todos los datos actuales por los del Excel "
                               "de ejemplo.\n\n¿Continuar?", parent=self):
            self.t.cargar_ejemplo()
            self.refrescar()

    def _exportar_csv(self) -> None:
        ruta = filedialog.asksaveasfilename(parent=self, defaultextension=".csv", initialfile="registros.csv",
                                            filetypes=[("CSV", "*.csv")])
        if ruta:
            self.t.exportar_csv(ruta)
            messagebox.showinfo("Exportado", f"Guardado en:\n{ruta}", parent=self)

    def _exportar_graficos(self) -> None:
        carpeta = filedialog.askdirectory(parent=self, title="Carpeta para los gráficos")
        if carpeta:
            archivos = Graficador(self.t).generar(carpeta)
            messagebox.showinfo("Gráficos", "Se crearon:\n" + "\n".join(a.name for a in archivos), parent=self)

    def _importar_excel(self) -> None:
        ruta = filedialog.askopenfilename(parent=self, title="Libro de Excel", filetypes=[("Excel", "*.xlsx")])
        if ruta and messagebox.askyesno("Importar", "Esto REEMPLAZA los datos actuales por los del Excel.\n\n"
                                        "¿Continuar?", parent=self):
            PuenteExcel.importar(self.t, ruta)
            self.t.guardar()
            self.refrescar()

    def _comparar_excel(self) -> None:
        ruta = filedialog.askopenfilename(parent=self, title="Libro de Excel a comparar",
                                          filetypes=[("Excel", "*.xlsx")])
        if not ruta:
            return
        copia = Tracker(RepositorioMemoria(), autoguardar=False)
        PuenteExcel.importar(copia, ruta)
        texto, _ = PuenteExcel.comparar(copia, ruta)
        ventana = tk.Toplevel(self)
        ventana.title("Comparación Excel vs Python")
        ventana.geometry("900x560")
        area = tk.Text(ventana, font=("Consolas", 10), wrap="none")
        area.insert("1.0", texto)
        area.config(state="disabled")
        area.pack(fill="both", expand=True)

    # ------------------------------------------------------------ pestaña: ayuda
    def _pestana_ayuda(self) -> None:
        f = ttk.Frame(self.nb, padding=12)
        self.nb.add(f, text="  Ayuda  ")
        area = tk.Text(f, font=(FUENTE, 11), wrap="word", bg=BLANCO, relief="flat", padx=16, pady=12)
        area.insert("1.0", AYUDA_GUI)
        area.config(state="disabled")
        area.pack(fill="both", expand=True)

    # ------------------------------------------------------------ refresco general
    def refrescar(self) -> None:
        self.lbl_cab.config(text=self._texto_cabecera())
        self._ajustar_combo(self.cb_cal, self._nombres(True))
        self._ajustar_combo(self.cb_mvr, self._nombres(False))
        self._refrescar_panel()
        self._refrescar_calendario()
        self._refrescar_habitos()
        self._refrescar_estadisticas()
        self._refrescar_modelo()
        self.var_nombre.set(self.t.usuario.nombre)
        self.var_meta.set(self.t.usuario.meta_general)
        self.lbl_inicio.config(text=f"Inicio del sistema: {ftexto_fecha(self.t.usuario.fecha_inicio)}")
        m = self.t.modelo
        for escala, valor in ((self.esc_amax, m.a_max), (self.esc_dias, m.dias_95), (self.esc_eps, m.epsilon)):
            if abs(escala.get() - valor) > 1e-9:
                escala.set(valor)


def lanzar_interfaz(tracker: Tracker, primera_vez: bool = False) -> None:
    app = AplicacionGrafica(tracker, primera_vez)
    app.mainloop()


# ======================================================================
# 12. Autopruebas (python tracker_habitos.py --probar)
# ======================================================================


def autopruebas() -> bool:
    resultados: List[Tuple[str, bool]] = []

    def comprobar(nombre: str, condicion: bool) -> None:
        resultados.append((nombre, bool(condicion)))
        print(("  OK       " if condicion else "  FALLÓ    ") + nombre)

    def cerca(a: float, b: float, tol: float = 1e-6) -> bool:
        return abs(a - b) < tol

    t = Tracker(RepositorioMemoria())
    t.fijar_hoy(date(2026, 9, 28))
    t.cargar_ejemplo()
    rep = t.reporte()
    e = {x.habito.nombre: x for x in rep.estadisticas()}
    leer, agua, ejer = e["Leer 20 minutos"], e["Tomar 2 litros de agua"], e["Hacer ejercicio"]
    print("Datos de ejemplo del Excel (90 registros)")
    comprobar("90 registros y 3 hábitos", len(t.registros) == 90 and len(t.habitos) == 3)
    comprobar("Rachas actuales 3 / 2 / 6 (igual que el Excel)",
              (leer.racha_actual, agua.racha_actual, ejer.racha_actual) == (3, 2, 6))
    comprobar("Mejores rachas 8 / 3 / 6", (leer.mejor_racha, agua.mejor_racha, ejer.mejor_racha) == (8, 3, 6))
    comprobar("% total 90 % / 76,67 % / 73,33 %",
              cerca(leer.pct_total, 0.9) and cerca(agua.pct_total, 23 / 30) and cerca(ejer.pct_total, 22 / 30))
    comprobar("Mejor racha global = 8 y promedio = 80 %",
              rep.mejor_racha_global() == 8 and cerca(rep.promedio_cumplimiento(), 0.8))

    print("Modelo de Cálculo")
    m = t.modelo
    comprobar("k = ln(20)/66 = 0,045389883", cerca(m.k, 0.0453898829326362, 1e-12))
    comprobar("t* = 33,3269", cerca(m.t_estrella, 33.3269012077348, 1e-9))
    comprobar("A(66) = 95 % del máximo", cerca(m.automaticidad(66), 95.0, 1e-9))
    comprobar("A'(0) = A_máx·k = 4,5390", cerca(m.velocidad(0), 4.53898829326362, 1e-9))
    comprobar("A'(t*) = ε = 1 (la definición de t*)", cerca(m.velocidad(m.t_estrella), 1.0, 1e-9))
    comprobar("A'' = −k·A' y siempre negativa",
              all(cerca(m.aceleracion(x), -m.k * m.velocidad(x), 1e-12) and m.aceleracion(x) < 0
                  for x in (0, 10, 33, 66, 90)))
    comprobar("Derivada numérica de A coincide con A'(t)",
              cerca((m.automaticidad(30 + 1e-6) - m.automaticidad(30 - 1e-6)) / 2e-6, m.velocidad(30), 1e-6))

    print("Modelo vs Real (hábito «Leer 20 minutos»)")
    res = t.modelo_vs_real("Leer 20 minutos").resultado()
    comprobar("T = 30, n = 27, ritmo = 0,9",
              res.dias_transcurridos == 30 and res.repeticiones == 27 and cerca(res.ritmo, 0.9))
    comprobar("A(27) = 70,6397 %", cerca(res.automaticidad_hoy, 70.6397021982305, 1e-9))
    comprobar("Faltan 39 repeticiones y 44 días", res.repeticiones_faltantes == 39 and res.dias_restantes == 44)
    comprobar("Fecha de consolidación 11/11/2026", res.fecha_consolidacion == date(2026, 11, 11))

    print("Reglas de negocio")
    desordenado = Tracker(RepositorioMemoria())
    desordenado.fijar_hoy(date(2026, 9, 28))
    h = desordenado.crear_habito("Meditar", "Bienestar")
    for dia in (3, 1, 2, 5, 4):                                    # registrados en desorden
        desordenado.registrar(h.id, True, date(2026, 9, dia))
    desordenado.registrar(h.id, False, date(2026, 9, 6))
    desordenado.registrar(h.id, True, date(2026, 9, 7))
    comprobar("La racha se calcula por fecha, sin importar el orden de registro",
              [r.racha for r in desordenado.registros_de(h)] == [1, 2, 3, 4, 5, 0, 1])
    try:
        desordenado.registrar(h.id, True, date(2026, 9, 7))
        comprobar("Rechaza un registro duplicado", False)
    except ErrorValidacion:
        comprobar("Rechaza un registro duplicado", True)
    try:
        desordenado.registrar(h.id, True, date(2026, 12, 1))
        comprobar("Rechaza fechas futuras", False)
    except ErrorValidacion:
        comprobar("Rechaza fechas futuras", True)
    try:
        desordenado.crear_habito("meditar")
        comprobar("Rechaza hábitos duplicados", False)
    except ErrorValidacion:
        comprobar("Rechaza hábitos duplicados", True)
    desordenado.editar_habito(h.id, nombre="Meditar 10 min")
    comprobar("Renombrar un hábito conserva su historial", len(desordenado.registros_de(h)) == 7)
    desordenado.archivar_habito(h.id)
    try:
        desordenado.registrar(h.id, True, date(2026, 9, 8))
        comprobar("Un hábito archivado no admite registros nuevos", False)
    except ErrorValidacion:
        comprobar("Un hábito archivado no admite registros nuevos", True)
    comprobar("…pero conserva su historial", len(desordenado.registros_de(h)) == 7)

    print("Calendario y escenarios")
    cal = Tracker(RepositorioMemoria())
    cal.fijar_hoy(date(2026, 9, 28))
    hh = cal.crear_habito("Estirar")
    d = date(2026, 9, 20)
    ciclo = [cal.alternar_dia(hh.id, d), cal.alternar_dia(hh.id, d), cal.alternar_dia(hh.id, d)]
    comprobar("Un clic en el calendario cicla: Sí -> No -> sin registro", ciclo == [True, False, None])
    comprobar("Eliminar un registro recalcula las rachas",
              (cal.registrar(hh.id, True, date(2026, 9, 21)), cal.registrar(hh.id, True, date(2026, 9, 22)),
               cal.eliminar_registro(hh.id, date(2026, 9, 21)))[2]
              and [r.racha for r in cal.registros_de(hh)] == [1])
    esc = t.modelo_vs_real("Leer 20 minutos").resultado(1.0)
    comprobar("Escenario «cumplo el 100 %» -> 39 días restantes", esc.dias_restantes == 39)

    print("Persistencia")
    copia = Tracker(RepositorioMemoria())
    copia.cargar_ejemplo()
    copia.guardar()
    otra = Tracker(copia.repositorio)
    comprobar("Guardar y cargar devuelve los mismos datos", otra.cargar() and otra.a_dict() == copia.a_dict())

    fallos = [n for n, ok in resultados if not ok]
    print(f"\n{len(resultados) - len(fallos)} de {len(resultados)} pruebas correctas.")
    return not fallos


# ======================================================================
# 13. Punto de entrada
# ======================================================================


def main(argv: Optional[Sequence[str]] = None) -> int:
    try:                                  # para que las barras y tildes salgan bien en Windows
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, OSError):
        pass

    parser = argparse.ArgumentParser(description="Tracker de Hábitos Inteligente (Python, POO)")
    parser.add_argument("--archivo", help="archivo JSON de datos (por defecto tracker_datos.json junto al programa)")
    parser.add_argument("--hoy", help="fija la fecha de hoy (AAAA-MM-DD); útil para comparar con el Excel")
    parser.add_argument("--consola", action="store_true", help="usa el menú de texto en vez de la ventana gráfica")
    parser.add_argument("--demo", action="store_true", help="muestra todo con los datos de ejemplo y sale")
    parser.add_argument("--probar", action="store_true", help="ejecuta las autopruebas y sale")
    parser.add_argument("--comparar", metavar="XLSX", help="compara los resultados con los del libro de Excel")
    parser.add_argument("--importar-excel", metavar="XLSX", help="trae hábitos y registros del Excel a este programa")
    parser.add_argument("--graficos", metavar="CARPETA", help="genera los gráficos PNG y sale")
    args = parser.parse_args(argv)

    try:
        if args.probar:
            return 0 if autopruebas() else 1

        ruta = Path(args.archivo) if args.archivo else Path(__file__).resolve().with_name("tracker_datos.json")
        usar_disco = not (args.demo or args.comparar)
        tracker = Tracker(RepositorioJSON(ruta) if usar_disco else RepositorioMemoria(),
                          autoguardar=usar_disco)
        modo_grafico = HAY_TK and not args.consola and not args.graficos and not args.importar_excel
        primera_vez = False
        if args.hoy:
            tracker.fijar_hoy(parsear_fecha(args.hoy))

        if args.demo:
            tracker.cargar_ejemplo()
            partes = [texto_dashboard(tracker), texto_habitos(tracker), texto_estadisticas(tracker),
                      texto_calculo(tracker.modelo), texto_modelo_vs_real(tracker.modelo_vs_real("Leer 20 minutos"))]
            print("\n\n".join(partes))
            return 0

        if args.comparar:
            PuenteExcel.importar(tracker, args.comparar)
            texto, ok = PuenteExcel.comparar(tracker, args.comparar)
            print(texto)
            return 0 if ok else 1

        if args.importar_excel:
            PuenteExcel.importar(tracker, args.importar_excel)
            tracker.guardar()
            print(f"Se importaron {len(tracker.habitos)} hábitos y {len(tracker.registros)} registros "
                  f"de {args.importar_excel}.")
        elif not tracker.cargar():
            if modo_grafico:
                primera_vez = True
            else:
                print("No hay datos guardados todavía.")
                resp = input("¿Cargar los datos de ejemplo del Excel (3 hábitos x 30 días)? [s/n]: ").strip().lower() \
                    if sys.stdin.isatty() else "s"
                if resp in ("s", "si", "sí"):
                    tracker.cargar_ejemplo()
                else:
                    tracker.usuario = Usuario(input("Tu nombre: ").strip() or "Estudiante")
                    tracker.guardar()

        if args.graficos:
            for archivo in Graficador(tracker).generar(args.graficos):
                print("Gráfico guardado:", archivo)
            return 0

        if modo_grafico:
            try:
                lanzar_interfaz(tracker, primera_vez)
                return 0
            except tk.TclError as error:
                print(f"No se pudo abrir la ventana ({error}). Se usa la versión de texto.")
        elif not args.consola and not HAY_TK:
            print("Este Python no trae Tkinter: se usa la versión de texto.")
        if primera_vez:
            tracker.cargar_ejemplo()
        MenuConsola(tracker).ejecutar()
        return 0
    except ErrorTracker as error:
        print(f"⚠ {error}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
